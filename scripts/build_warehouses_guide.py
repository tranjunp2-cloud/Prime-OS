"""Build the Vietnamese team guide for the current Warehouses prototype."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output/docs/prime-os-cam-nang-warehouses.docx"
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(0.65)
section.left_margin = section.right_margin = Inches(0.72)
section.header_distance = section.footer_distance = Inches(0.28)

styles = doc.styles
for border in styles.element.xpath(".//w:pBdr"):
    border.getparent().remove(border)
for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
    styles[name].font.name = "Arial"
    styles[name].font.color.rgb = RGBColor.from_string("000000")
normal = styles["Normal"]
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.12
normal.paragraph_format.space_after = Pt(6)
for name, size, before, after in [("Title", 26, 0, 9), ("Heading 1", 20, 0, 9), ("Heading 2", 13, 12, 6), ("Heading 3", 11, 8, 4)]:
    style = styles[name]
    style.font.size = Pt(size)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True
styles["Subtitle"].font.size = Pt(12)
styles["Subtitle"].paragraph_format.space_after = Pt(9)

header = section.header.paragraphs[0]
header.text = "PRIME OS  |  CẨM NANG VẬN HÀNH"
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor.from_string("596579")
footer = section.footer.paragraphs[0]
footer.paragraph_format.space_after = Pt(0)
footer.text = "Warehouses · UI prototype · 30/09/2026"
footer.paragraph_format.tab_stops.add_tab_stop(Inches(6.70))
footer.add_run("\t")
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)
for run in footer.runs:
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string("596579")

def para(text, label=None):
    p = doc.add_paragraph()
    if label:
        p.add_run(label + (" " if label in ("•", "1.", "2.", "3.") else ": ")).bold = True
    p.add_run(text)
    return p

def h(text):
    return doc.add_heading(text, level=2)

def step(number, text):
    p = para(text, f"{number}.")
    p.paragraph_format.left_indent = Inches(0.20)
    p.paragraph_format.first_line_indent = Inches(-0.20)
    return p

def bullet(text):
    p = para(text, "•")
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    return p

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    props = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            e.set(qn(f"w:{key}"), value)
        borders.append(e)
    props.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:w"), "95" if edge in ("top", "bottom") else "115")
        e.set(qn("w:type"), "dxa")
        margins.append(e)
    props.append(margins)
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i, header_text in enumerate(headers):
        t.rows[0].cells[i].text = header_text
    for values in rows:
        cells = t.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = value
    for ri, row in enumerate(t.rows):
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "E9EDF3")
                cell._tc.get_or_add_tcPr().append(shading)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.08
                for run in p.runs:
                    run.font.size = Pt(10.5)
                    run.font.bold = ri == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

def new_page(title):
    doc.add_page_break()
    doc.add_heading(title, level=1)

# Page 1 — mental model and reading the current UI.
doc.add_heading("Cẩm nang Warehouses", 0)
doc.add_paragraph("Quản lý kho và tồn hàng trên Prime OS", "Subtitle")
para("Dành cho team sản phẩm, thiết kế và vận hành. Tài liệu mô tả UI prototype hiện tại; tên nút giữ tiếng Anh để dễ tìm trên màn hình. Các phần còn mô phỏng được ghi rõ ở cuối tài liệu.")
h("1 Nguyên tắc cần nhớ")
para("Thay đổi tồn tại kho sẽ cập nhật tồn trong Product Master. Tồn trên từng sàn mặc định độc lập, chỉ gửi cập nhật khi người dùng chủ động thiết lập cho listing.")
table(["Loại kho", "Có được sửa tồn không"], [
    ("Kho nội bộ và kho 3PL có quyền quản lý tồn trong Prime OS", "Có. Điều chỉnh và chuyển kho theo từng sản phẩm hoặc SKU biến thể."),
    ("Kho do sàn quản lý như FBA hoặc FBS", "Không. Chỉ xem số liệu do bên quản lý kho ghi nhận."),
    ("Kho ảo hoặc vị trí chưa xác định", "Không. Cần xác định lại loại kho và quyền quản lý trước."),
], [2.80, 4.26])
para("Không phải kho nào đồng bộ từ sàn về cũng có thể sửa. Quyền quản lý tồn quyết định khả năng thao tác, không chỉ việc kho đó là kho vật lý.", "Lưu ý")
h("2 Tra cứu tồn kho")
step(1, "Vào Warehouses → My warehouses. Dùng All warehouses để xem tổng hoặc chọn một kho trong Choose one warehouse.")
step(2, "Tìm theo tên sản phẩm hoặc SKU. Mở rộng dòng sản phẩm để xem tồn theo kho hoặc từng biến thể.")
step(3, "Dùng bộ lọc trạng thái tồn, kênh bán và sắp xếp số lượng để tìm sản phẩm cần xử lý.")
table(["Hiển thị", "Cách hiểu"], [
    ("0 · Out of stock", "Đã ghi nhận số lượng và hiện không còn hàng."),
    ("— · Not recorded", "Chưa ghi nhận số lượng; không được hiểu là 0."),
    ("Partial data", "Một số biến thể chưa có số liệu. Tổng mới gồm phần đã ghi nhận."),
], [2.00, 5.06])
para("Total across warehouses cộng số đã ghi nhận ở các kho, kể cả kho ngoài. Đây không phải toàn bộ số hàng Prime OS được quyền bán. Bộ lọc Selling on chỉ thể hiện kênh có listing đang hoạt động, không chứng minh đã bật đồng bộ tồn.", "Đọc đúng số liệu")

# Page 2 — warehouse creation and first stock count.
new_page("Tạo kho và nhập tồn ban đầu")
h("3 Tạo kho nội bộ")
para("Warehouses → Add warehouse → Add Physical Warehouse.", "Vào đâu")
step(1, "Nhập Warehouse Name, Warehouse Code và địa chỉ. Tên kho, mã kho và Street Address là các trường tối thiểu để mở nút tạo.")
step(2, "Kiểm tra thông tin và bấm Create Physical Warehouse. Mã kho được chuyển thành chữ in hoa.")
step(3, "Quay lại danh sách và chọn kho mới để kiểm tra. Tạo kho chưa làm phát sinh số lượng tồn và chưa tự liên kết hoặc đồng bộ tồn với shop.")
para("Có thêm một kho nội bộ để ghi nhận hàng. Form này chưa phải luồng tạo kho FBA/FBS hay cấu hình kho 3PL.", "Kết quả")
h("4 Ghi nhận tồn lần đầu cho sản phẩm")
para("Products → Product Master → mở sản phẩm → Pricing & Inventory → Inventory by location → Add stock location.", "Vào đâu")
step(1, "Chọn kho đang hoạt động, có quyền sửa và chưa ghi nhận tồn cho sản phẩm này.")
step(2, "Nhập Initial stock bằng số lượng thực tế. Có thể nhập 0 nếu đã kiểm đếm và kho đang trống.")
step(3, "Chọn Adjustment reason, kiểm tra số lượng rồi bấm Save location & stock. Bấm Cancel nếu chưa muốn lưu.")
para("Kho và Product Master có số tồn đầu tiên; lịch sử ghi Initial count. Tồn trên sàn không thay đổi.", "Kết quả")
table(["Tình huống", "Thao tác đúng"], [
    ("Chưa từng ghi nhận tồn tại kho", "Dùng Add stock location để ghi số đầu tiên; không dùng New Adjustment."),
    ("Đã có số tồn, kể cả 0", "Dùng Adjust stock để cập nhật số lượng thực tế."),
    ("Sản phẩm có nhiều biến thể", "Điều chỉnh riêng từng Variant SKU đã có tồn. Luồng thêm vị trí mới cho biến thể chưa hoàn thiện."),
], [2.75, 4.31])
para("Kho HCM mới tạo chưa có số liệu cho sản phẩm A. Ghi Initial stock = 100 sẽ tạo bản ghi “Not recorded → 100”, không coi đây là điều chỉnh từ 0.", "Ví dụ")
para("Add stock location hiện hỗ trợ sản phẩm Single. Kho FBA/FBS không xuất hiện như một lựa chọn được phép nhập tồn thủ công.", "Phạm vi hiện tại")

# Page 3 — the two stock mutations and their audit trail.
new_page("Điều chỉnh và chuyển kho")
h("5 Điều chỉnh theo số kiểm đếm")
para("Warehouses → chọn kho và sản phẩm → Adjust stock. Hoặc Stock actions & history → New Adjustment.", "Vào đâu")
step(1, "Chọn Warehouse và Product. Nếu có biến thể, chọn đúng Variant SKU.")
step(2, "Xem Current stock, nhập New stock là tổng số thực tế còn lại, không phải số muốn cộng hoặc trừ.")
step(3, "Chọn Adjustment reason, kiểm tra chênh lệch rồi bấm Record adjustment.")
para("Hệ thống đang ghi 100, kiểm đếm còn 92 → nhập 92. Hệ thống tự tính giảm 8, lưu lý do và cập nhật Product Master.", "Ví dụ")
para("Chỉ nhận số nguyên từ 0 trở lên, có thay đổi và có lý do. Nếu tồn bị thay đổi trong lúc mở form, đóng rồi mở lại để kiểm tra số mới. SKU chưa có số tồn hoặc kho chỉ đọc sẽ bị chặn.", "Kiểm tra trước khi lưu")
h("6 Chuyển hàng giữa hai kho")
para("Warehouses → Stock actions & history → Transfer Stock.", "Vào đâu")
step(1, "Chọn Source Warehouse và Destination Warehouse. Hai kho phải khác nhau và đều có quyền sửa tồn.")
step(2, "Chọn Product or SKU; với sản phẩm nhiều biến thể, chọn SKU cụ thể.")
step(3, "Nhập Transfer Quantity là số nguyên dương, không vượt quá số đã ghi nhận tại kho nguồn. Bấm Create Transfer.")
para("Chuyển 20 từ HCM có 100 sang Hà Nội có 30 → HCM còn 80, Hà Nội có 50. Tổng hai kho vẫn là 130; tồn trên sàn không đổi.", "Kết quả")
para("Prototype trừ kho nguồn và cộng kho đích ngay khi xác nhận. Chưa có luồng thực tế xuất kho → đang vận chuyển → nhận hàng; cũng chưa kiểm tra lượng đã giữ cho đơn trước khi chuyển.", "Giới hạn quan trọng")
h("7 Kiểm tra lịch sử")
para("Mở Stock actions & history → Stock Adjustments để xem SKU, kho, số trước và sau, lý do, thời gian. Lần ghi nhận đầu tiên hiển thị Initial count; thao tác điều chỉnh hiển thị phần tăng hoặc giảm.")
para("Lịch sử điều chỉnh được lưu trên thiết bị hiện tại; Operator đang là Local demo. Danh sách Stock Transfers vẫn là dữ liệu mẫu, chưa tự thêm phiếu từ thao tác chuyển kho.", "Phân biệt")

# Page 4 — optional external stock publication and honest prototype boundaries.
new_page("Liên kết gian hàng và phạm vi prototype")
h("8 Liên kết shop khác với đồng bộ tồn")
para("Tại Warehouses, mở Locations & shops hoặc Warehouse details → Linked shops để xem các shop liên quan và trạng thái kết nối. Manage shop connections đưa bạn tới Sales Channels → Connected Channels.")
para("Một shop đã liên kết với kho không có nghĩa mọi thay đổi tồn đều được gửi lên shop đó. Muốn gửi tồn, cần cấu hình riêng cho listing.")
para("Product Master → mở sản phẩm → Sales Channels → Manage listing → Price & inventory → Adjust inventory → Stock sync policy.", "Cấu hình ở đâu")
table(["Lựa chọn", "Ý nghĩa"], [
    ("Do not publish stock", "Không gửi số tồn. Giữ tồn của listing độc lập với Master."),
    ("Automatic sync", "Thiết lập lấy tồn từ nguồn đã chọn; có thể giữ lại Safety buffer và giới hạn bằng Channel allocation cap."),
    ("Manual update", "Nhập Quantity to send cho listing, không sửa tồn thực tế tại kho."),
], [2.10, 4.96])
para("Xem số lượng dự kiến trước khi Review & publish. Kết quả trong prototype chưa xác nhận tồn đã cập nhật thật lên sàn. Với Amazon FBA, tồn do Amazon quản lý; Master không ghi đè.")
h("9 Tồn hiện có và ATP không phải một số")
para("Stock on hand là lượng đã ghi nhận tại kho. ATP là lượng có thể cam kết cho đơn sau khi trừ phần đã giữ và mức dự phòng. Không dùng Total across warehouses thay cho ATP.")
para("Stock breakdown trong phần nâng cao đang dùng số mẫu. ATP được minh họa bằng On hand − Reserved − Safety stock; nếu kết quả âm thì hiển thị 0. Chưa liên thông đầy đủ với đơn hàng và việc giữ hoặc nhả tồn.")
h("10 Những phần cần hoàn thiện tiếp")
bullet("Tạo kho: dữ liệu kho mới chỉ giữ trong phiên hiện tại. Manager, Phone Number và Allow Negative Stock chưa được lưu hoặc áp dụng đầy đủ.")
bullet("Tồn biến thể: cần bổ sung luồng nhập số đầu tiên cho SKU chưa có tồn tại một vị trí mới.")
bullet("Chuyển kho: cần phiếu chuyển thực, trạng thái giao nhận và lịch sử. Hiện mới thay đổi số lượng ngay trong prototype.")
bullet("Mapping và routing: Link Warehouse hiện chỉ báo thao tác; các tùy chọn chọn kho giao hàng chưa điều khiển đơn thực tế. Đồng bộ với nhà cung cấp cần được kiểm chứng khi tích hợp.")
para("Chọn đúng kho và SKU → kiểm tra số đã ghi nhận → lưu số thực tế kèm lý do → kiểm tra Product Master và lịch sử. Chỉ thiết lập gửi tồn ra sàn khi có nhu cầu rõ ràng.", "Tóm tắt cho team")

doc.core_properties.title = "Cẩm nang Warehouses"
doc.core_properties.subject = "Hướng dẫn luồng kho và tồn hàng theo UI prototype Prime OS"
doc.core_properties.author = "Prime OS"
doc.core_properties.keywords = "Warehouses, tồn kho, Product Master, ATP, hướng dẫn"
doc.core_properties.language = "vi-VN"
lang = OxmlElement("w:lang")
lang.set(qn("w:val"), "vi-VN")
styles["Normal"]._element.get_or_add_rPr().append(lang)
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
