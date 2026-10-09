"""Create the Vietnamese seller guide from the Orders UI reviewed on 2026-10-09.

Evidence: Orders.tsx, OrderWorkspaceDetail.tsx, ManualOrderDialog.tsx,
OrderPreparationDialog.tsx, OrderProcessingDialog.tsx, OrderBatchDialog.tsx,
OrderShipmentJourney.tsx, order-work-queues.ts, order-detail-flow.ts,
order-prototype.ts, orders-api.ts and packages/order-workflow/rules.js.
Run with the bundled document runtime; render and inspect before delivery.
"""
from pathlib import Path
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output/docs/prime-os-cam-nang-seller-orders.docx"
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(0.65)
section.left_margin = section.right_margin = Inches(0.72)
section.header_distance = section.footer_distance = Inches(0.27)

for border in doc.styles.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet'):
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:hAnsi'), 'Arial')
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.12
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.widow_control = True
for name, size, before, after in [('Title', 25, 0, 8), ('Heading 1', 20, 0, 12), ('Heading 2', 13, 13, 6)]:
    style = doc.styles[name]
    style.font.size = Pt(size)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True
doc.styles['Subtitle'].font.size = Pt(12)
doc.styles['Subtitle'].paragraph_format.space_after = Pt(10)

# A small running identifier and page number help when sellers share excerpts.
header = section.header.paragraphs[0]
header.text = 'PRIME OS   |   HƯỚNG DẪN SELLER'
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor(0, 0, 0)
footer = section.footer.paragraphs[0]
footer.text = 'Orders  ·  09/10/2026'
footer.paragraph_format.tab_stops.add_tab_stop(Inches(6.72))
footer.add_run('\t')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
for run in footer.runs:
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string('595959')


def rich(p, text):
    for i, part in enumerate(re.split(r'\*\*(.*?)\*\*', text)):
        run = p.add_run(part)
        run.bold = bool(i % 2)
    return p


def para(text):
    return rich(doc.add_paragraph(), text)


def step(number, text):
    p = para(f'**{number}.** {text}')
    p.paragraph_format.left_indent = Inches(0.20)
    p.paragraph_format.first_line_indent = Inches(-0.20)
    return p


def bullet(text):
    p = rich(doc.add_paragraph(style='List Bullet'), text)
    p.paragraph_format.space_after = Pt(5)
    return p


def h(text):
    doc.add_heading(text, level=2)


def page(title):
    p = doc.add_heading(title, level=1)
    p.paragraph_format.page_break_before = True


def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    props = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}')
        for k, v in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            e.set(qn(f'w:{k}'), v)
        borders.append(e)
    props.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for edge in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:w'), '100' if edge in ('top', 'bottom') else '125')
        e.set(qn('w:type'), 'dxa')
        margins.append(e)
    props.append(margins)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for i, text in enumerate(headers):
        rich(t.rows[0].cells[i].paragraphs[0], text)
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            rich(cells[i].paragraphs[0], text)
    for ri, row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'E8EEF5')
                cell._tc.get_or_add_tcPr().append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.08
                for run in p.runs:
                    run.font.size = Pt(10.5)
                    if ri == 0:
                        run.bold = True
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1
    p.add_run().font.size = Pt(3)
    return t


# Page 1 — distinguish order progress, payment and fulfillment ownership.
doc.add_heading('Cẩm nang Orders dành cho Seller', 0)
doc.add_paragraph('Quản lý đơn hàng trên Prime OS', 'Subtitle')
para('Hướng dẫn tra cứu đơn, xác nhận thanh toán, chuẩn bị hàng, theo dõi giao nhận và xử lý phát sinh. Tên nút giữ bằng tiếng Anh để bạn tìm đúng trên màn hình.')
para('**Phạm vi sử dụng:** Đơn có nhãn **Prototype** chỉ dùng để thử luồng; không thu tiền, đặt vận chuyển hay thay đổi tồn thật. Với đơn từ sàn chưa kết nối thao tác, tiếp tục xử lý tại **Seller Center** của sàn.')
h('1 Tìm đúng đơn cần xử lý')
para('Vào **Main Workspace → Orders**. Tìm theo mã đơn, khách hàng, số điện thoại, SKU hoặc mã vận đơn; kết hợp bộ lọc **Shop**, **Payment** và **Order status**. Mở đơn để xem thông tin và hành động tiếp theo.')
table(['Nhóm hoặc trạng thái', 'Ý nghĩa và việc cần làm'], [
    ('**Draft orders**', 'Bản nháp chưa gửi xử lý; nằm riêng, không tính trong All.'),
    ('**Awaiting confirmation**', 'Kiểm tra khách, địa chỉ, sản phẩm và điều kiện thanh toán trước khi xác nhận.'),
    ('**To ship**', 'Gồm Preparing là đang chuẩn bị và Ready to ship là đã đóng gói, chờ giao cho hãng vận chuyển.'),
    ('**Shipping**', 'Hàng đã bàn giao; theo dõi hành trình của từng kiện.'),
    ('**Delivered / Completed**', 'Delivered là đã giao. Completed là đã đóng luồng đơn; tiền sàn chuyển về được theo dõi riêng.'),
    ('**Returns & refunds / Cancelled**', 'Tra cứu đơn có phát sinh trả hàng hoặc hoàn tiền; xem riêng đơn đã hủy.'),
], [2.10, 4.96])
h('2 Xác định ai chuẩn bị và đóng gói')
para('**Seller fulfilled:** Bạn hoặc kho của bạn soạn, đóng gói và bàn giao hàng. **Marketplace fulfilled:** Kho của sàn thực hiện, bạn theo dõi cập nhật từ sàn.')
para('**Seller shipping / Platform shipping** cho biết bên thu xếp vận chuyển. Dùng hãng vận chuyển của sàn không có nghĩa sàn đóng gói thay bạn.')
para('**Cần nhớ:** Trạng thái đơn, trạng thái thanh toán và tiền thực nhận là ba thông tin khác nhau. Không đánh dấu đã thu tiền chỉ vì đơn đã giao.')

# Page 2 — intake and payment have real side-effect boundaries.
page('Tiếp nhận và xác nhận đơn')
h('1 Kiểm tra đơn có sẵn')
step(1, 'Mở **Awaiting confirmation**, chọn đúng shop và đơn cần xử lý. Kiểm tra mã đơn, người nhận, địa chỉ, SKU, số lượng và tổng tiền.')
step(2, 'Xem **Payment**, **Dispatch warehouse** và **Ship-by deadline**. Mốc Updated trên danh sách là lúc tải dữ liệu trang, không phải thời điểm đồng bộ sàn.')
step(3, 'Nếu có **Needs review**, đọc lý do và xử lý trước. Dùng **Assign** để giao người phụ trách, **Edit note** để ghi chú nội bộ khi có quyền.')
step(4, 'Bấm **Confirm Order** khi nút được phép thao tác. Sau xác nhận, đơn chuyển sang **Preparing** trong nhóm **To ship**.')
para('**Kết quả cần kiểm tra:** Đơn đã xác nhận, đúng thông tin giao hàng và không còn vướng mắc ngăn xử lý. Xác nhận đơn chưa giữ hoặc trừ tồn kho.')
h('2 Tạo đơn thủ công')
para('**Vào đâu:** Orders → **Create order**. Cần quyền ghi đơn hàng.')
step(1, 'Nhập **Order reference**, **Sales channel / Store**, tiền tệ, khách hàng và địa chỉ. Email và mã bưu chính là thông tin tùy chọn trên form.')
step(2, 'Thêm SKU, tên, số lượng và đơn giá ở **Order items**; kiểm tra Discount, Shipping fee, Tax và Order total. Form hiện nhập SKU trực tiếp, chưa phải bước xác thực tồn.')
step(3, 'Chọn **Payment method**, **Payment status**, **Fulfillment type**, **Shipping arranged by** và kho ưu tiên. Chỉ chọn Paid khi đã có bằng chứng thanh toán.')
step(4, 'Chọn **Save draft** để giữ bản nháp hoặc **Create order** để đưa vào hàng chờ xác nhận. Với bản nháp, dùng **Edit draft** hoặc **Submit draft order → Submit order**.')
h('3 Ghi nhận thanh toán đúng thời điểm')
para('Mở **Payment details → Confirm payment**, nhập **Payment reference / Evidence**, rồi bấm **Record payment** sau khi đã đối chiếu khoản tiền thực nhận. Thao tác này ghi nhận thanh toán có sẵn, không trừ tiền thẻ của khách.')
table(['Trường hợp', 'Cách xử lý'], [
    ('Trả trước chưa Paid', 'Cần ghi nhận khoản đã nhận trước khi chuẩn bị hàng. Một số đơn có bằng chứng chờ kiểm tra sẽ bị chặn ngay ở bước xác nhận.'),
    ('COD còn Unpaid', 'Có thể chuẩn bị và giao hàng. Chỉ ghi Paid sau khi xác nhận đã thu được tiền.'),
    ('Đơn nguồn sàn hoặc chỉ có quyền xem', 'Theo hướng dẫn trên màn hình; xử lý tại kênh nguồn hoặc nhờ người có quyền, không tìm cách bỏ qua nút bị khóa.'),
], [2.10, 4.96])

# Page 3 — the manual order path is distinct from the compact demo dialog.
page('Chuẩn bị và giao đơn thủ công')
para('Áp dụng cho đơn **Seller fulfilled**, vận chuyển do seller thu xếp và tài khoản có quyền xử lý. Đây là luồng tách từng bước; đơn mẫu dùng form gộp ở trang tiếp theo.')
h('1 Chọn kho và kiểm tra hàng')
para('Từ đơn đã xác nhận, bấm **Prepare order**. Điền **Dispatch warehouse**, thêm **Ship-by deadline** nếu cần và kiểm tra hàng thực tế trước khi lưu. Đơn bắt đầu chuẩn bị; thao tác này không tự giữ tồn.')
h('2 Ghi nhận vận đơn đã có')
para('Liên hệ hãng vận chuyển để thu xếp giao nhận trước, sau đó bấm **Add carrier shipment**. Nhập **Carrier**, **Service** và **Carrier-issued tracking number**; có thể thêm liên kết nhãn, cách bàn giao và thời gian đã hẹn.')
para('**Lưu ý:** Lưu mã vận đơn trong Prime OS không tự đặt lịch lấy hàng hay tạo nhãn của hãng. Dùng nhãn gốc do đơn vị vận chuyển cấp.')
h('3 Kiểm hàng và đóng gói')
step(1, 'Bấm **Pack & verify**. Nhập số đã kiểm cho từng dòng hoặc dùng **Scan or enter SKU → Verify unit** để kiểm từng đơn vị.')
step(2, 'Đảm bảo số đã kiểm khớp toàn bộ số lượng trong đơn. Nếu cùng SKU xuất hiện nhiều dòng, kiểm riêng từng dòng.')
step(3, 'Nhập cân nặng và kích thước kiện; xác nhận đã kiểm nội dung và dán nhãn. Bấm **Finish packing**.')
para('**Kết quả:** Đơn sang **Ready to ship**, vẫn trong nhóm To ship. Đóng gói xong chưa có nghĩa hãng vận chuyển đã nhận hàng.')
h('4 Ghi nhận bàn giao cho hãng vận chuyển')
para('Chỉ bấm **Record carrier pickup** sau khi hãng đã nhận kiện. Điền **Carrier pickup evidence / Reference**, xác nhận hãng đã tiếp nhận rồi lưu. Đơn chuyển sang **Shipping**.')
h('5 Ghi nhận giao thành công')
para('Theo dõi **Shipment journey** theo mã vận đơn. Khi có bằng chứng giao nhận, dùng **Record delivery**, nhập **Delivery evidence / Reference** rồi lưu. Trạng thái chuyển sang **Delivered**, thanh toán không tự đổi thành Paid.')
para('**Hai lưu ý:** **Print packing slip** là phiếu soạn/đóng gói, không thay nhãn vận chuyển. Với **Platform shipping** chưa có kết nối thao tác, thu xếp giao hàng trong Seller Center theo thông báo trên UI.')

# Page 4 — prototype walkthrough, marketplace path and safe completion.
page('Thử luồng trên đơn mẫu và hoàn tất đơn')
h('1 Đơn mẫu do seller xử lý')
para('Đơn có nhãn **Prototype** dùng form **Prepare shipment** gộp kiểm tra thanh toán, kho, vận chuyển và đóng gói. Kết quả chỉ lưu trong tab trình duyệt hiện tại.')
step(1, 'Sau **Confirm Order**, mở **Prepare shipment**. Kiểm tra thông tin điền sẵn; dùng **Edit** ở phần cần điều chỉnh.')
step(2, 'Nếu trả trước chưa nhận tiền, điền bằng chứng và xác nhận thanh toán mô phỏng. Với thẻ, dùng **Simulate payment update**; COD có thể tiếp tục khi chưa Paid.')
step(3, 'Kiểm tra kho xuất và vận chuyển. Nếu seller tự thu xếp, bổ sung hãng, dịch vụ và mã vận đơn. Nếu sàn thu xếp, form mô phỏng dữ liệu do sàn cấp.')
step(4, 'Tại **Verify items**, nhập đủ số đã kiểm hoặc quét SKU. Kích thước kiện có thể để trống trong form mẫu; nếu nhập thì điền đủ và hợp lệ.')
step(5, 'Chọn **Parcel packed and shipping label attached**, rồi **Mark ready to ship**. Sau đó thử **Record carrier pickup** để chuyển sang Shipping.')
para('**Use sample values** chỉ điền dữ liệu thử nhanh. **Restart prototype** đưa đơn mẫu về điểm bắt đầu để thử lại, không phải thao tác khôi phục đơn thật.')
h('2 Đơn mẫu do kho sàn xử lý')
para('Với **Marketplace fulfilled**, seller không tự soạn và đóng gói. Dùng **Simulate marketplace update** để thử trạng thái sẵn sàng giao, rồi **Simulate carrier pickup** để thử bàn giao. Các nút này không gửi lệnh tới sàn.')
h('3 Theo dõi hành trình và hoàn tất')
para('Trong **Shipment journey**, nút **Simulate tracking update** cho phép thử các cập nhật vận chuyển được phép ở bước hiện tại, gồm giao thành công hoặc thất bại. Dòng **Simulated tracking** cho biết chưa có kết nối hãng vận chuyển thật.')
para('Với đơn cho phép hoàn tất, bấm **Complete order** khi đã Delivered, đã Paid, tất cả kiện đã giao và không còn yêu cầu trả hàng hoặc vấn đề đang mở. Đơn chuyển sang **Completed**; thay đổi vòng đời đơn bị khóa.')
para('**Shop settlement** là thông tin tiền shop nhận sau phí và điều chỉnh. **Estimated net settlement** chỉ là ước tính; Completed không có nghĩa sàn đã chuyển tiền về.')
h('4 Đọc đúng thông tin tồn')
para('**Inventory snapshot / Sample stock** là số ghi nhận tại thời điểm kiểm tra, không phải tồn trực tiếp. **Not checked** là chưa kiểm tra, không phải hết hàng. Không dùng thao tác tạo, xác nhận hoặc chuẩn bị đơn làm bằng chứng đã giữ hay trừ tồn.')

# Page 5 — exceptions and efficient day-to-day operation.
page('Xử lý phát sinh và nhiều đơn cùng lúc')
h('1 Tạm dừng hoặc hủy đơn')
para('Khi cần làm rõ thông tin trước bàn giao, chọn **More → Hold Order**, nhập lý do rồi **Put on hold**. Khi đã giải quyết, chọn **Release hold** và nhập lý do; đơn tiếp tục ở bước cũ.')
para('Chỉ dùng **Cancel order** khi nút còn được phép ở giai đoạn sớm. Nhập lý do và xác nhận. **Hủy đơn không tự hoàn tiền**; nếu đơn đã Paid, cần xử lý khoản hoàn riêng. Không dùng hủy đơn để thay cho quy trình trả hàng sau giao.')
h('2 Giao thất bại và trả hàng')
step(1, 'Tìm đơn ở **Shipping → Delivery failed** hoặc **Returns & refunds**. Kiểm tra mã vận đơn, diễn biến và liên hệ người phụ trách để xác nhận hướng xử lý.')
step(2, 'Mở **Review return** hoặc **Returns / RMA** để xem mã yêu cầu, trạng thái và lý do. RMA là hồ sơ xử lý hàng trả.')
step(3, 'Đối chiếu riêng ba việc: hãng báo hàng quay về, kho kiểm nhận hàng thực tế, và khoản hoàn tiền. Không coi trạng thái Returned là đã nhập lại tồn hoặc đã hoàn tiền.')
para('**Phạm vi hiện tại:** Orders chủ yếu hiển thị và theo dõi hồ sơ trả hàng. Chưa có luồng đầy đủ ngay tại đây để tạo yêu cầu đổi trả, duyệt hoàn tiền hoặc tự nhập lại tồn. Đơn sàn cần xử lý tại Seller Center theo quyền và hướng dẫn của kênh.')
h('3 Xác nhận và in phiếu hàng loạt')
step(1, 'Lọc đúng shop và trạng thái. Chọn các đơn bằng checkbox, rồi **Confirm orders** hoặc **Print packing slips**.')
step(2, 'Muốn áp dụng cho toàn bộ kết quả lọc, chọn **Available actions → Can confirm** hoặc **Can print packing slips**, rồi dùng nút xử lý tất cả. Phạm vi có thể gồm nhiều trang, không chỉ trang đang xem.')
step(3, 'Kiểm tra danh sách đủ điều kiện và đơn bị bỏ qua trước khi chạy. Sau xử lý, đọc kết quả từng đơn; xử lý hoặc thử lại riêng phần thất bại, không làm lại đơn đã thành công.')
para('Cho phép pop-up nếu bị chặn khi in. Mở hộp thoại in chưa xác nhận giấy đã in. **Export orders** xuất các đơn theo bộ lọc; **Export selected** chỉ xuất các đơn đã chọn.')
h('4 Khi không thể tiếp tục')
table(['Thông báo hoặc tình huống', 'Cần làm gì'], [
    ('**Needs review / On hold**', 'Đọc lý do, xử lý vướng mắc và nhả tạm dừng trước khi tiếp tục.'),
    ('**Reload latest order**', 'Tải bản mới rồi kiểm tra lại; người khác có thể vừa cập nhật đơn.'),
    ('**Read-only / Channel actions not connected**', 'Nhờ người có quyền hoặc xử lý ở kênh nguồn theo thông báo.'),
    ('Nhiều kiện hoặc giao một phần', 'Xem từng kiện riêng; luồng xử lý đơn một kiện không áp dụng cho toàn bộ đơn.'),
], [2.60, 4.46])

doc.core_properties.title = 'Cẩm nang Orders dành cho Seller'
doc.core_properties.subject = 'Hướng dẫn quản lý đơn hàng trên Prime OS'
doc.core_properties.author = 'Prime OS'
doc.core_properties.keywords = 'Orders, Seller, Prime OS, tiếng Việt'
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
