from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "reports" / "primeos-feature-overview-uiux-optimization-report-2026-08-17.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = "172033"
INDIGO = "5B4CF0"
INDIGO_DARK = "3F35C9"
BLUE = "2563EB"
TEAL = "0F766E"
GREEN = "15803D"
AMBER = "B45309"
RED = "B42318"
INK = "171717"
BODY = "33363F"
MUTED = "667085"
LIGHT = "F6F7FB"
LAVENDER = "EFEDFF"
PALE_BLUE = "EAF4FF"
PALE_GREEN = "EAF8F1"
BORDER = "D9DDE7"
WHITE = "FFFFFF"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def configure_table(table, widths, header=True):
    total = sum(widths)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row_i, row in enumerate(table.rows):
        if header and row_i == 0:
            tr_pr = row._tr.get_or_add_trPr()
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            tr_pr.append(repeat)
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    set_font(run, size=8.7, color=BODY)
        if header and row_i == 0:
            for cell in row.cells:
                shade(cell, NAVY)
                for p in cell.paragraphs:
                    for run in p.runs:
                        set_font(run, size=8.5, color=WHITE, bold=True)


def set_font(run, name="Aptos", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_field(run, field):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def paragraph_border_bottom(paragraph, color=INDIGO, size=12, space=8):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.78)
sec.bottom_margin = Inches(0.72)
sec.left_margin = Inches(0.86)
sec.right_margin = Inches(0.86)
sec.header_distance = Inches(0.35)
sec.footer_distance = Inches(0.35)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10.3)
normal.font.color.rgb = RGBColor.from_string(BODY)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.1

for name, size, color, before, after in (
    ("Title", 30, NAVY, 0, 8),
    ("Heading 1", 17, NAVY, 16, 8),
    ("Heading 2", 13.5, INDIGO_DARK, 12, 6),
    ("Heading 3", 11.5, NAVY, 9, 4),
):
    st = styles[name]
    st.font.name = "Aptos Display" if name != "Normal" else "Aptos"
    st._element.rPr.rFonts.set(qn("w:ascii"), st.font.name)
    st._element.rPr.rFonts.set(qn("w:hAnsi"), st.font.name)
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for style_name in ("List Bullet", "List Number"):
    st = styles[style_name]
    st.font.name = "Aptos"
    st.font.size = Pt(10.2)
    st.paragraph_format.left_indent = Inches(0.5)
    st.paragraph_format.first_line_indent = Inches(-0.25)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.1

caption = styles["Caption"]
caption.font.name = "Aptos"
caption.font.size = Pt(8.5)
caption.font.italic = True
caption.font.color.rgb = RGBColor.from_string(MUTED)
caption.paragraph_format.space_before = Pt(3)
caption.paragraph_format.space_after = Pt(8)

if "Callout" not in [s.name for s in styles]:
    callout = styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
else:
    callout = styles["Callout"]
callout.font.name = "Aptos"
callout.font.size = Pt(10.5)
callout.font.bold = True
callout.font.color.rgb = RGBColor.from_string(NAVY)
callout.paragraph_format.left_indent = Inches(0.18)
callout.paragraph_format.right_indent = Inches(0.18)
callout.paragraph_format.space_before = Pt(7)
callout.paragraph_format.space_after = Pt(9)

header = sec.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = hp.add_run("PRIMEOS  /  PRODUCT & UX ASSESSMENT")
set_font(r, size=8, color=MUTED, bold=True)
paragraph_border_bottom(hp, color=BORDER, size=4, space=4)

footer = sec.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = fp.add_run("Internal working report  |  17 Aug 2026   •   ")
set_font(r, size=8, color=MUTED)
r = fp.add_run()
set_font(r, size=8, color=MUTED)
add_field(r, "PAGE")


def add_para(text="", *, size=None, color=None, bold=False, italic=False, align=None, before=0, after=6, style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def add_bullet(text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    if level:
        p.paragraph_format.left_indent = Inches(0.75)
    p.add_run(text)
    return p


def add_number(text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def add_callout(label, text, fill=LAVENDER):
    table = doc.add_table(rows=1, cols=1)
    configure_table(table, [9360], header=False)
    cell = table.cell(0, 0)
    shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label.upper() + "  ")
    set_font(r, size=9, color=INDIGO_DARK, bold=True)
    r = p.add_run(text)
    set_font(r, size=10.2, color=NAVY, bold=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_table(headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    configure_table(table, widths, header=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_picture(path, width=6.68, caption_text=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(path), width=Inches(width))
    alt_text = caption_text or f"Ảnh minh họa giao diện PrimeOS: {Path(path).stem}"
    shape._inline.docPr.set("descr", alt_text)
    shape._inline.docPr.set("title", Path(path).stem)
    if caption_text:
        cp = doc.add_paragraph(caption_text, style="Caption")
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER


def page_break():
    doc.add_page_break()


# Cover — named override: editorial report cover within standard_business_brief.
add_para("PRIME COMMERCE", size=10, color=INDIGO, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=82, after=20)
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(10)
r = p.add_run("PrimeOS")
set_font(r, name="Aptos Display", size=34, color=NAVY, bold=True)
add_para("BÁO CÁO TỔNG QUAN TÍNH NĂNG\n& ĐỀ XUẤT TỐI ƯU UI/UX", size=20, color=NAVY, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
rule = doc.add_paragraph()
paragraph_border_bottom(rule, color=INDIGO, size=18, space=4)
add_para("Đánh giá hiện trạng prototype • Định hướng trải nghiệm vận hành • Roadmap 90 ngày", size=12, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, before=10, after=54)
add_callout("Product thesis", "Demand tạo cơ hội. Customer giữ ngữ cảnh. COS thực thi. Intelligence học và tối ưu. Finance chuyển commerce evidence thành funding readiness.", fill=LIGHT)
add_para("Phiên bản 1.0", size=10, color=MUTED, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=48, after=2)
add_para("Ngày đánh giá: 17/08/2026", size=10, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
add_para("Phạm vi: apps/web, apps/admin, apps/api và tài liệu sản phẩm hiện hành", size=9, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)

page_break()
doc.add_heading("Tóm tắt điều hành", level=1)
add_callout("Kết luận", "PrimeOS đã có nền tảng sản phẩm mạnh và khác biệt ở closed-loop commerce. Vấn đề lớn nhất hiện nay không phải thiếu tính năng, mà là mật độ thông tin, độ sâu điều hướng và tính nhất quán giữa các surface làm tăng chi phí nhận thức của operator.")
add_para("PrimeOS hiện là một prototype commerce operating system kết nối Demand, CRM/Customer, E-commerce/COS, Intelligence, Finance và Prime AI trong cùng một vòng vận hành. Codebase thể hiện quy mô đáng kể: khoảng 111 route declarations, 61 page components và 193 component TSX trong web workspace. Hệ thống đã có design token, component primitives, light/dark theme, i18n ba ngôn ngữ và các mẫu operating loop dùng chung.")
doc.add_heading("Đánh giá nhanh", level=2)
add_table(
    ["Trụ cột", "Hiện trạng", "Đánh giá", "Ưu tiên"],
    [
        ("Giá trị sản phẩm", "Closed-loop từ signal → decision → execution → outcome", "Mạnh", "Giữ và làm rõ"),
        ("Kiến trúc thông tin", "Nhiều Area/Tower/Floor, route cũ/mới song song", "Cần tối giản", "P0"),
        ("Desktop operations", "Data-dense, phù hợp operator nhưng lặp nhiều panel", "Khá", "P1"),
        ("Mobile", "Icon rail + nội dung hẹp; bảng/filter chưa chuyển mô hình", "Yếu", "P0"),
        ("Accessibility", "Có Radix primitives nhưng chưa có bằng chứng audit toàn diện", "Chưa đủ", "P0"),
        ("Performance", "Bundle web chính khoảng 2.1 MB minified; route import eager", "Rủi ro", "P1"),
        ("Design system", "Token và component nền đã có, tài liệu còn lệch đường dẫn", "Khá", "P1"),
        ("Quality", "Build pass; test baseline còn lỗi navigation/i18n/copilot/theme", "Cần ổn định", "P0"),
    ],
    [1700, 3900, 1500, 2260],
)
doc.add_heading("Ba quyết định nên chốt ngay", level=2)
add_number("Chốt IA theo 5 khu vực vận hành chính và tách Platform Admin khỏi navigation nghiệp vụ.")
add_number("Thiết kế responsive theo nhiệm vụ, không thu nhỏ nguyên desktop; mobile tập trung queue, search, approve và detail.")
add_number("Thiết lập UX quality gate: WCAG AA, keyboard, responsive matrix, route-state integrity và performance budget.")

doc.add_heading("Phạm vi và phương pháp", level=1)
add_para("Báo cáo kết hợp desk review tài liệu, đọc cấu trúc code, kiểm tra route/navigation, đối chiếu screenshot release-proof và áp dụng heuristic cho enterprise SaaS/data-dense dashboard. Đây là đánh giá product/UX ở mức hệ thống, không thay thế usability test với người dùng thật.")
add_table(
    ["Nguồn", "Nội dung sử dụng"],
    [
        ("System description & maps", "Product thesis, domain map, data architecture, limitations, route inventory"),
        ("Source code", "PrimeRoutes, navigation tree, component/page counts, design token usage"),
        ("Release-proof screenshots", "Desktop hierarchy, density, action placement, mobile adaptation"),
        ("UI/UX heuristics", "Accessibility, interaction, performance, responsive, forms, navigation, data visualization"),
        ("Build/test evidence", "Build success, lint warnings, current failing test baseline"),
    ],
    [2600, 6760],
)
doc.add_heading("1. Tổng quan hệ thống và giá trị sản phẩm", level=1)
doc.add_heading("1.1 Định vị", level=2)
add_para("PrimeOS không phải một CRM, OMS hay dashboard đơn lẻ. Hệ thống đóng vai trò lớp vận hành và intelligence bao quanh commerce execution, giúp operator đi từ tín hiệu thị trường đến quyết định, handoff và kết quả có thể truy vết.")
add_callout("North Star UX", "Mỗi màn hình phải trả lời nhanh bốn câu hỏi: Điều gì đang xảy ra? Tại sao quan trọng? Ai cần hành động? Kết quả sẽ được đọc lại ở đâu?")
doc.add_heading("1.2 Kiến trúc dịch vụ", level=2)
add_table(
    ["Service", "Vai trò", "Người dùng chính", "Ranh giới hiện tại"],
    [
        ("Web operator", "Không gian vận hành đa domain", "Operator, manager", "React/Vite; nhiều store frontend và fallback data"),
        ("Admin control room", "Quản trị curated data, IAM và cấu hình", "Admin", "UI tập trung trong App.tsx lớn; file-backed API"),
        ("Express API", "Auth, IAM, Growth OS, task, finance/feedback", "Web/Admin", "Prototype persistence, chưa phải production SoT"),
        ("Prime AI", "Đọc ngữ cảnh, giải thích, đề xuất, tạo draft", "Operator", "Deterministic; không tự ghi business state"),
    ],
    [1700, 2700, 1900, 3060],
)
doc.add_heading("1.3 Vòng vận hành cốt lõi", level=2)
for label, text in (
    ("Signal", "Demand, marketplace, social, VOC, inventory, customer và finance cung cấp tín hiệu."),
    ("Decision", "Intelligence/Prime AI tổng hợp evidence và đề xuất next-best action."),
    ("Handoff", "Công việc được chuyển đúng tower/owner với SLA và guardrail."),
    ("Execution", "COS/CRM/Service/Finance thực thi nghiệp vụ trong ranh giới quyền hạn."),
    ("Outcome", "Order, RFQ, fulfillment, service và campaign result quay lại làm evidence."),
):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(label + " — ")
    set_font(r, bold=True, color=INDIGO_DARK)
    p.add_run(text)

doc.add_heading("2. Tổng quan tính năng theo khu vực", level=1)
feature_rows = [
    ("Performance", "KPI, funnel, journey loop, leads, connectors, AI actions", "Điều hành tổng quan và phát hiện điểm nghẽn", "MVP mạnh"),
    ("Intelligence", "Decision Hub, signals, launch decisions, operation/branding/consulting agent", "Biến evidence thành quyết định có kiểm soát", "Khác biệt lõi"),
    ("Ecom / COS", "Product, inventory, orders, fulfillment, policies, event & audit", "Execution source of truth của commerce", "Lõi vận hành"),
    ("CRM / Demand", "MDEC, sources, campaigns, content, leads/RFQs, re-engage", "Thu nhận và chuyển đổi nhu cầu đa kênh", "Phạm vi rộng"),
    ("Customer", "Profile, identity matching, timeline, service/RMA/SLA", "Giữ ngữ cảnh khách hàng xuyên hành trình", "Cần sâu hơn"),
    ("Finance", "Readiness, evidence, document, lender route, risk blockers", "Chuẩn hóa funding readiness", "Định vị đúng"),
    ("Platform Admin", "Account, workspace, role, member, invitation, audit", "Quản trị tenant và quyền", "Prototype"),
    ("Prime AI", "Global copilot, draft actions, contextual prompts, approval boundary", "Giảm thao tác và tăng tốc quyết định", "Cần trust UX"),
]
add_table(["Khu vực", "Tính năng chính", "Giá trị", "Mức trưởng thành"], feature_rows, [1500, 3700, 2700, 1460])
doc.add_heading("3. Phân tích trải nghiệm hiện tại", level=1)
doc.add_heading("3.1 Điểm mạnh", level=2)
for text in (
    "Ngôn ngữ operating loop (Signal → Decision → Handoff → Outcome) tạo mental model tốt cho hệ thống phức tạp.",
    "Navigation dùng icon Lucide và label, có active state rõ; command search giúp truy cập nhanh entity và màn hình.",
    "Design token đã bao phủ surface, typography, status, area, spacing, sidebar/header và motion.",
    "Các primitive Radix/shadcn hỗ trợ nền tảng accessibility tốt hơn custom control thuần túy.",
    "Route legacy được giữ để giảm gián đoạn thói quen người dùng và deep link cũ.",
    "Bộ screenshot đa ngôn ngữ và đa viewport là nền tảng QA tốt.",
):
    add_bullet(text)

img1 = ROOT / "research/screenshots/phase-7-release-proof/01-intelligence-decision-hub-desktop-en.png"
add_picture(img1, 6.68, "Hình 1. Decision Hub thể hiện rõ operating loop và handoff, nhưng phần above-the-fold sử dụng nhiều khối có trọng lượng thị giác tương đương.")

doc.add_heading("3.2 Vấn đề chính trên desktop", level=2)
add_table(
    ["Vấn đề", "Bằng chứng", "Tác động", "Đề xuất"],
    [
        ("Trùng lặp thông tin", "Hero, KPI, chips, operating loop và handoff cùng lặp area/tower/readiness", "Tăng thời gian quét", "Giảm 25–35% nội dung above-the-fold"),
        ("Action hierarchy", "Open handoff, Back to loop, Open links và Prime AI cạnh tranh", "Không rõ primary action", "Mỗi page chỉ một primary CTA theo trạng thái"),
        ("Navigation sâu", "Area → tower → floor + query/hash/redirect", "Mất định hướng và state", "Breadcrumb chuẩn + canonical URL + state preservation"),
        ("Card density", "Nhiều card viền mảnh, mức nhấn gần nhau", "Khó phân biệt summary và work queue", "Dùng surface hierarchy 3 cấp"),
        ("Status semantics", "Ready/readiness/alerts dùng nhiều badge và màu", "Nguy cơ color-only", "Luôn kèm icon, label và mô tả nguyên nhân"),
    ],
    [1800, 2500, 2200, 2860],
)

img2 = ROOT / "research/screenshots/phase-7-release-proof/05-ecom-oms-detail-desktop-en.png"
add_picture(img2, 6.68, "Hình 2. Order detail có hierarchy nghiệp vụ tốt; nên làm rõ action safety, timeline semantics và trạng thái liên domain.")
doc.add_heading("Cải tiến nhanh cho Order Detail", level=3)
for text in (
    "Nhóm action theo intent: primary fulfillment, secondary routing và destructive cancel; giải thích điều kiện khả dụng.",
    "Hiển thị impact preview trước khi gửi fulfillment: warehouse, carrier, SLA, inventory reservation và affected records.",
    "Timeline dùng icon + verb + actor + timestamp nhất quán; cho phép lọc System/User/Connector và mở audit detail.",
    "Đưa exception lên đầu trang khi SLA, inventory hoặc payment có rủi ro; trạng thái bình thường không nên chiếm cùng mức ưu tiên.",
):
    add_bullet(text)

page_break()
doc.add_heading("3.3 Vấn đề responsive/mobile", level=2)
img3 = ROOT / "research/screenshots/phase-7-release-proof/04-customer-account-profile-mobile-en.png"
add_picture(img3, 3.0, "Hình 3. Mobile hiện giữ rail icon và layout desktop thu nhỏ, khiến vùng nội dung hẹp và filter stack dài.")
add_callout("P0 mobile", "Không nên xem mobile là bản desktop thu nhỏ. Với operator mobile, ưu tiên 5 tác vụ: tìm kiếm, xem queue, mở detail, approve/reject và nhận alert.", fill=PALE_BLUE)
for text in (
    "Thay rail icon cố định bằng top app bar + navigation drawer hoặc bottom navigation tối đa 5 điểm đến.",
    "Chuyển filter thành bottom sheet; hiển thị số filter đang áp dụng và nút Clear all.",
    "Biến data table thành list card có progressive disclosure; giữ action chính trong sticky bottom bar.",
    "Bảo đảm touch target tối thiểu 44×44 px, khoảng cách 8 px và không dựa vào hover.",
    "Giữ state filter/search/scroll khi back; tránh redirect làm reset context.",
):
    add_bullet(text)

doc.add_heading("3.4 Accessibility và trust", level=2)
add_table(
    ["Tiêu chí", "Chuẩn mục tiêu", "Rủi ro cần kiểm tra"],
    [
        ("Contrast", "WCAG AA: 4.5:1 text, 3:1 large/UI", "Muted text, pastel status, dark/light parity"),
        ("Keyboard", "Mọi thao tác, dialog, table và chart dùng được bằng bàn phím", "Dense nav, command palette, custom table"),
        ("Focus", "Focus ring 2–4 px; focus main sau đổi route", "Sticky shell và modal/drawer"),
        ("Forms", "Label thật, validation on blur, lỗi cạnh field và recovery", "Create/edit flows, integrations, finance documents"),
        ("AI trust", "Evidence, confidence, permission, preview và audit", "Recommendation/action có thể bị hiểu là tự động"),
        ("Data viz", "Không dùng màu đơn độc; có table/text summary", "KPI, funnel, readiness, status charts"),
    ],
    [1900, 3800, 3660],
)

doc.add_heading("4. Định hướng UI/UX mục tiêu", level=1)
doc.add_heading("4.1 Design principles", level=2)
principles = [
    ("Action before decoration", "Mọi màn hình ưu tiên công việc phải làm tiếp theo, không ưu tiên số lượng card."),
    ("Progressive disclosure", "Tóm tắt trước, evidence và configuration mở dần khi cần."),
    ("One canonical path", "Một destination có một URL chuẩn; redirect chỉ phục vụ tương thích."),
    ("Evidence-first AI", "Mọi gợi ý AI hiển thị nguồn, độ tin cậy, tác động và ranh giới quyền."),
    ("Responsive by task", "Desktop tối ưu điều phối; mobile tối ưu triage và approval."),
    ("Accessible by default", "Component definition bao gồm keyboard, focus, contrast, empty/loading/error."),
]
for title, body in principles:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(title + ": ")
    set_font(r, bold=True, color=INDIGO_DARK)
    p.add_run(body)

doc.add_heading("4.2 Information architecture đề xuất", level=2)
add_table(
    ["Cấp 1", "Cấp 2 điển hình", "Nguyên tắc"],
    [
        ("Overview", "Performance, My Work, Alerts", "Điểm vào cá nhân hóa theo role"),
        ("Intelligence", "Decisions, Signals, Agents", "Không trộn insight và execution list"),
        ("Commerce", "Products, Inventory, Orders, Fulfillment", "COS là execution core"),
        ("Growth", "Sources, Campaigns, Content, Leads/RFQs", "Gộp các hành trình demand liên tục"),
        ("Customers", "Profiles, Service, Re-engage", "Bám entity và lifecycle"),
        ("Finance", "Readiness, Evidence, Applications", "Ngôn ngữ không overclaim"),
        ("Platform", "Team, Integrations, Automation, Audit", "Tách khỏi daily operations"),
    ],
    [1700, 3600, 4060],
)

page_break()
doc.add_heading("5. Đề xuất tối ưu theo nhóm", level=1)
doc.add_heading("5.1 Shell và navigation", level=2)
for text in (
    "Giới hạn sidebar mặc định ở 6–7 nhóm; nhóm ít dùng chuyển vào Platform hoặc More.",
    "Hiển thị breadcrumb đầy đủ từ cấp sâu thứ ba; tên không bị truncate ở desktop rộng.",
    "Dùng canonical route registry làm nguồn duy nhất cho route, navigation, breadcrumb, permission và i18n label.",
    "Command palette ưu tiên recent, role-based actions và entity result; hỗ trợ keyboard shortcut được công bố rõ.",
    "Khi chuyển route, focus vào h1/main; khi back, khôi phục scroll và filter state.",
):
    add_bullet(text)

doc.add_heading("5.2 Dashboard và data density", level=2)
for text in (
    "Above-the-fold chỉ gồm: page intent, health summary, một primary action và exceptions cần xử lý.",
    "KPI card phải có baseline, trend, timeframe và click-through; tránh số liệu không có ngữ cảnh.",
    "Queue/table trên 50 dòng dùng virtualization; filter có saved view và density toggle.",
    "Loading trên 1 giây dùng skeleton có kích thước cố định để tránh CLS; lỗi có Retry và timestamp.",
    "Charts có direct labels/tooltip, text summary và export CSV; không dùng donut trên 5 category.",
):
    add_bullet(text)

doc.add_heading("5.3 Forms, workflow và feedback", level=2)
for text in (
    "Form dài chia section/step, autosave draft và cảnh báo khi đóng khi chưa lưu.",
    "Validation on blur; submit error focus vào field lỗi đầu tiên và có error summary nếu lỗi nhiều.",
    "Action async khóa double-submit, hiển thị progress và success/error có recovery path.",
    "Delete/cancel order cần confirm theo mức rủi ro; tác vụ có thể hoàn tác dùng Undo toast.",
    "Read-only khác disabled cả về semantics và visual; quyền thiếu phải giải thích thay vì ẩn im lặng.",
):
    add_bullet(text)

doc.add_heading("5.4 Prime AI interaction model", level=2)
add_table(
    ["Lớp", "UI bắt buộc", "Mục tiêu trust"],
    [
        ("Context", "Entity, timeframe, filter và role đang dùng", "Người dùng biết AI đang nhìn gì"),
        ("Evidence", "Nguồn dữ liệu, freshness, links", "Có thể kiểm chứng"),
        ("Recommendation", "Lý do, confidence, trade-off", "Không biến gợi ý thành sự thật"),
        ("Action preview", "Before/after, affected records, permission", "Ngăn mutation ngoài ý muốn"),
        ("Approval", "Approve/Edit/Reject và optional reason", "Human-in-the-loop rõ ràng"),
        ("Audit", "Actor, timestamp, inputs, decision, result", "Traceability"),
    ],
    [1500, 3900, 3960],
)

doc.add_heading("6. Design system mục tiêu", level=1)
add_para("Giữ hướng Genesis hiện tại: light-first, editorial precision, surface phẳng, indigo cho interaction. Không áp dụng palette AI tím-hồng một cách máy móc; PrimeOS nên dùng indigo hiện hữu làm brand accent và màu semantic theo domain/status.")
add_table(
    ["Token/Pattern", "Khuyến nghị"],
    [
        ("Typography", "General Sans heading; DM Sans body/UI; JetBrains Mono cho ID; body desktop ≥14px, mobile ≥16px"),
        ("Spacing", "4/8px base; component 8/12/16; section 24/32; page 32/48"),
        ("Radius", "Button/input 6px; panel 8px; card 12px; tránh radius tùy ý"),
        ("Elevation", "3 cấp: flat, raised, overlay; không dùng shadow cho mọi card"),
        ("Motion", "150–300ms; transform/opacity; hỗ trợ prefers-reduced-motion"),
        ("Status", "Icon + label + color; semantic mapping ổn định trong cả light/dark"),
        ("Density", "Comfortable/Compact modes; row 44/36px nhưng touch target vẫn ≥44px trên touch"),
        ("Breakpoints", "375, 768, 1024, 1440; test landscape và 200% zoom"),
    ],
    [2300, 7060],
)
doc.add_heading("7. Backlog ưu tiên", level=1)
backlog = [
    ("P0", "Canonical route + nav registry", "Route/navigation/i18n/breadcrumb", "Giảm drift và test lỗi", "M"),
    ("P0", "Responsive shell mobile", "Header, drawer/bottom nav, content gutter", "Tăng usable width", "M"),
    ("P0", "Accessibility baseline", "Contrast, keyboard, focus, skip link, labels", "WCAG AA critical paths", "M"),
    ("P0", "Stabilize failing tests", "Navigation, i18n, Copilot, theme", "Quality gate đáng tin", "S"),
    ("P1", "Dashboard hierarchy", "Overview, Decision, Leads/RFQs", "Giảm time-to-action", "L"),
    ("P1", "Route-level code splitting", "React.lazy/Suspense, heavy charts", "Giảm initial bundle", "M"),
    ("P1", "Unified data states", "Loading/empty/error/stale/offline", "Nhất quán feedback", "M"),
    ("P1", "Table/list responsive pattern", "Saved view, density, mobile card", "Tăng hiệu suất operator", "L"),
    ("P1", "AI evidence/action preview", "Sources, confidence, approval, audit", "Tăng trust", "L"),
    ("P2", "Personalized overview", "Role-based queue và recent work", "Giảm navigation", "L"),
    ("P2", "Design token lint", "Cấm raw color/spacing ngoài allowlist", "Giảm visual drift", "M"),
    ("P2", "Usability analytics", "Task success, abandonment, time-on-task", "Tối ưu dựa dữ liệu", "M"),
]
add_table(["Priority", "Hạng mục", "Phạm vi", "Kết quả", "Effort"], backlog, [850, 2400, 2500, 2750, 860])

doc.add_heading("8. Roadmap triển khai 90 ngày", level=1)
add_table(
    ["Giai đoạn", "Tuần", "Mục tiêu", "Deliverables"],
    [
        ("Foundation", "1–3", "Ổn định nền tảng UX", "Route registry; fix test baseline; a11y audit; responsive shell spec"),
        ("Core journeys", "4–7", "Tối ưu hành trình có giá trị cao", "Overview; Leads/RFQs; Order detail; Customer profile; AI approval"),
        ("Systemization", "8–10", "Chuẩn hóa component/state", "Data state; responsive table; forms; token lint; Storybook-like catalog"),
        ("Validation", "11–12", "Đo và khóa quality gate", "Usability test; performance budget; Playwright matrix; release report"),
    ],
    [1700, 900, 2600, 4160],
)
doc.add_heading("Đội hình tối thiểu", level=2)
add_bullet("01 Product/UX lead: IA, journey, prioritization, usability test.")
add_bullet("01 Product designer: responsive patterns, components, prototypes, accessibility specs.")
add_bullet("02 Frontend engineers: shell, route registry, performance, component refactor.")
add_bullet("01 QA/automation: keyboard, responsive, visual regression và test stabilization.")
add_bullet("Part-time domain owners: validate semantics cho COS, CRM, Finance và AI.")

doc.add_heading("9. KPI và tiêu chí thành công", level=1)
add_table(
    ["KPI", "Baseline", "Mục tiêu 90 ngày", "Cách đo"],
    [
        ("Task success rate", "Chưa đo", ">= 85% cho 5 critical tasks", "Moderated usability + event analytics"),
        ("Time to next action", "Chưa đo", "Giảm >= 25%", "Từ page load đến primary action"),
        ("Navigation depth", "Nhiều route/query/redirect", "<= 3 tương tác cho critical destination", "Journey instrumentation"),
        ("Accessibility", "Chưa có gate toàn diện", "0 critical/serious axe issue", "Automated + manual keyboard/SR"),
        ("Mobile usability", "Desktop thu nhỏ", ">= 80% task success", "375px/390px user test"),
        ("Initial JS", "~2.1 MB main chunk minified", "Giảm >= 35%", "Vite bundle report"),
        ("Test quality", "12 web failures quan sát", "0 known failures trong CI", "Vitest/Playwright"),
        ("UX consistency", "80 raw hex occurrences quan sát", "Raw value giảm, token adoption >= 95%", "Static lint"),
    ],
    [2100, 2300, 2600, 2360],
)

page_break()
doc.add_heading("10. Acceptance checklist", level=1)
checks = [
    "Mỗi route có h1 duy nhất, title, breadcrumb và canonical URL.",
    "Mọi thao tác đạt keyboard-only; focus visible và không có keyboard trap.",
    "Text contrast đạt 4.5:1; color không phải tín hiệu duy nhất.",
    "Mobile 375px không có document-level horizontal scroll.",
    "Touch target tối thiểu 44×44px; primary controls không phụ thuộc hover.",
    "Loading/empty/error/stale/success có pattern thống nhất và recovery path.",
    "Form có label, required state, validation on blur và focus first error.",
    "Destructive actions có confirm/undo phù hợp và mô tả impact.",
    "AI recommendation có evidence, confidence, preview, permission và audit.",
    "Chart có legend/label, summary văn bản và accessible data alternative.",
    "Route/page nặng được lazy-load; layout không nhảy khi data về.",
    "Test chạy sạch cho navigation, i18n, theme, Copilot và critical journeys.",
]
for item in checks:
    add_bullet("☐ " + item)

doc.add_heading("11. Rủi ro và phụ thuộc", level=1)
add_table(
    ["Rủi ro", "Tác động", "Kiểm soát"],
    [
        ("Tối ưu UI trước khi khóa IA", "Rework nhiều", "Chốt canonical taxonomy và role journeys trước component polish"),
        ("Refactor file lớn cùng lúc", "Regression logic", "Tách theo feature boundary, characterization test trước"),
        ("Dữ liệu prototype không phản ánh scale", "Sai density/performance assumption", "Test dataset 10× và latency simulation"),
        ("Thiếu user research", "Tối ưu theo ý kiến nội bộ", "5–8 operator tests theo role và task"),
        ("AI wording overclaim", "Mất trust/risk pháp lý", "UX copy guardrail và approval policy"),
    ],
    [2800, 2200, 4360],
)

doc.add_heading("12. Kết luận", level=1)
add_para("PrimeOS đã vượt qua giai đoạn “dashboard demo” nhờ có product thesis rõ, các domain liên kết và một operating loop có thể kể thành câu chuyện. Để chuyển thành sản phẩm operator-grade, hệ thống cần giảm độ sâu và sự lặp lại trong trải nghiệm, biến mobile thành một mode tác vụ riêng, chuẩn hóa trust UX cho AI và đưa accessibility/performance vào quality gate.")
add_callout("Khuyến nghị cuối", "Không mở rộng thêm diện tích tính năng trong 90 ngày tới nếu chưa ổn định 5 critical journeys: Overview → Decision, Lead/RFQ → Customer, Order → Fulfillment, Customer → Service và AI recommendation → Approved action.", fill=PALE_GREEN)

doc.add_heading("Phụ lục A — Nguồn tham chiếu nội bộ", level=1)
sources = [
    "docs/09-primeos-full-system-description.md",
    "docs/02-prime-os-system-map.md",
    "docs/03-screen-map.md",
    "docs/platform/prime-os-design-token-map.md",
    "docs/platform/prime-os-genesis-redesign-implementation.md",
    "apps/web/src/routes/PrimeRoutes.tsx",
    "apps/web/src/lib/prime/prime-navigation.ts",
    "apps/web/src/index.css",
    "research/screenshots/phase-7-release-proof/",
]
for source in sources:
    add_bullet(source)

doc.add_heading("Phụ lục B — Lưu ý về mức độ tin cậy", level=2)
add_para("Các số lượng route/component và bundle size là snapshot tại thời điểm đánh giá. Các đề xuất usability là heuristic kết hợp bằng chứng giao diện; cần được xác nhận bằng analytics và usability testing trước khi dùng làm cam kết KPI chính thức.")

# Keep headings with following content and prevent lonely captions.
for p in doc.paragraphs:
    if p.style and p.style.name.startswith("Heading"):
        p.paragraph_format.keep_with_next = True

doc.core_properties.title = "PrimeOS — Báo cáo tổng quan tính năng và tối ưu UI/UX"
doc.core_properties.subject = "Product feature overview, UX audit and 90-day optimization roadmap"
doc.core_properties.author = "PrimeOS Product & UX"
doc.core_properties.keywords = "PrimeOS, product overview, UI, UX, accessibility, roadmap"
doc.save(OUT)
print(OUT)
