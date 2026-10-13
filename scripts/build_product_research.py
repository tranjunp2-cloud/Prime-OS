from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output/docs/prime-os-product-research.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.65), Inches(.62)
sec.left_margin = sec.right_margin = Inches(.72)
sec.header_distance = sec.footer_distance = Inches(.28)

styles = doc.styles
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet', 'List Number']:
    styles[name].font.name = 'Arial'
    styles[name]._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Arial')
    styles[name].font.color.rgb = RGBColor(0, 0, 0)
normal = styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.10
normal.paragraph_format.space_after = Pt(6)
for name, size in [('Title', 29), ('Heading 1', 19), ('Heading 2', 12.5), ('Heading 3', 11)]:
    s = styles[name]
    s.font.size = Pt(size)
    s.font.bold = True
    s.paragraph_format.space_before = Pt(12 if name != 'Title' else 0)
    s.paragraph_format.space_after = Pt(7)
    s.paragraph_format.keep_with_next = True
styles['Subtitle'].font.size = Pt(12)
styles['Subtitle'].font.italic = False
styles['Subtitle'].paragraph_format.space_after = Pt(10)
for style in styles:
    for border in style._element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
for name in ['List Bullet', 'List Number']:
    styles[name].paragraph_format.space_after = Pt(5)
    styles[name].paragraph_format.line_spacing = 1.08
    styles[name].paragraph_format.left_indent = Inches(.18)
    styles[name].paragraph_format.first_line_indent = Inches(-.13)

# Low-key running furniture; no decorative title rules or boxed summaries.
h = sec.header.paragraphs[0]
h.text = 'PRIME OS   /   PRODUCT RESEARCH'
h.runs[0].font.size = Pt(8.5)
h.runs[0].font.color.rgb = RGBColor.from_string('626262')
f = sec.footer.paragraphs[0]
f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = f.add_run('09.10.2026   •   ')
r.font.size = Pt(9)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
f._p.append(field)
doc.core_properties.title = 'Product Master và quản lý listing đa sàn'
doc.core_properties.subject = 'Research và cơ sở quyết định Product, UX và kỹ thuật'
doc.core_properties.author = 'Prime OS'
doc.core_properties.keywords = 'Product Master, PIM, listing, mapping, SKU, readiness, sync, seller UX'

def p(text, style=None):
    return doc.add_paragraph(text, style)

def bullet(text):
    return p(text, 'List Bullet')

def head(text):
    return doc.add_heading(text, 2)

def page(title, strap=None):
    heading = doc.add_heading(title, 1)
    heading.paragraph_format.page_break_before = True
    if strap:
        p(strap)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    tPr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        e = OxmlElement('w:' + edge)
        for k, v in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            e.set(qn('w:' + k), v)
        borders.append(e)
    tPr.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for edge in ['top', 'left', 'bottom', 'right']:
        e = OxmlElement('w:' + edge)
        e.set(qn('w:w'), '95')
        e.set(qn('w:type'), 'dxa')
        margins.append(e)
    tPr.append(margins)
    for i, width in enumerate(widths):
        t.columns[i].width = Inches(width)
    for i, label in enumerate(headers):
        t.rows[0].cells[i].text = label
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    repeat = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for ridx, row in enumerate(t.rows):
        cant = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(cant)
        for i, cell in enumerate(row.cells):
            cell.width = Inches(widths[i])
            if ridx == 0 or ridx % 2 == 0:
                sh = OxmlElement('w:shd')
                sh.set(qn('w:fill'), '262B33' if ridx == 0 else 'F5F6F7')
                cell._tc.get_or_add_tcPr().append(sh)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(2)
                para.paragraph_format.space_before = Pt(0)
                para.paragraph_format.line_spacing = 1.04
                for run in para.runs:
                    run.font.size = Pt(10.3)
                    if ridx == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
    p('').paragraph_format.space_after = Pt(0)
    return t

def link(para, label, url):
    rel_id = para.part.relate_to(url, RT.HYPERLINK, is_external=True)
    node = OxmlElement('w:hyperlink')
    node.set(qn('r:id'), rel_id)
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '245A81')
    props.append(color)
    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')
    props.append(underline)
    run.append(props)
    text = OxmlElement('w:t')
    text.text = label
    run.append(text)
    node.append(run)
    para._p.append(node)


# 1 — Decision brief, not a decorative cover.
doc.add_heading('Product Master và\nquản lý listing đa sàn', 0)
p('Research và cơ sở quyết định cho Product · Design · Engineering', 'Subtitle')
p('Prime OS  |  Ngày 09 tháng 10 năm 2026  |  Phiên bản 1.0')
head('Kết luận để trình bày với team')
p('Hướng tối ưu là cho seller liên kết và xử lý dữ liệu ngay trong luồng review, nhưng tách rõ việc lưu quan hệ, kích hoạt Master và đồng bộ từng listing. Không buộc seller hoàn thiện mọi yêu cầu của mọi sàn trước khi quản lý sản phẩm tập trung; cũng không coi “đã link” là “đã sẵn sàng sync”.')
head('Năm quyết định nền tảng')
bullet('Nhiều listing có thể thuộc một Master khi seller xác nhận cùng sản phẩm. Chọn nhiều dòng không mặc nhiên có nghĩa là gộp thành một sản phẩm.')
bullet('Master Active dựa trên dữ liệu lõi SKU, tên, mô tả và giá hợp lệ. Field bắt buộc riêng của sàn chỉ chặn sync listing tương ứng, không chặn toàn bộ Master.')
bullet('Cho phép “Link now, finish later” với trạng thái cần xử lý. Cặp SKU chưa xác nhận không được dùng để sync giá hoặc tồn.')
bullet('Hiển thị nguồn dữ liệu, field đích và giá trị sẽ lưu; sửa ngay tại bảng review. AI hoặc quy tắc chỉ đưa ra đề xuất, không thay seller quyết định ghi đè.')
bullet('Liên kết, Active và lưu Master không tự publish, không tự bật sync, không tự cộng hoặc chia tồn kho.')
head('Cơ sở và giới hạn bằng chứng')
p('Tài liệu tổng hợp các tình huống và quyết định trong trao đổi ngày 06–09/10/2026; đối chiếu mã prototype hiện tại; nghiên cứu tài liệu chính thức của Akeneo, Shopify, Amazon, NN/g và W3C. Đây là desk research kết hợp đánh giá thiết kế, chưa phải kết quả phỏng vấn seller hay kiểm chứng hệ thống production.')
p('Các nguồn [R1–R8] ở cuối tài liệu hỗ trợ nguyên tắc thiết kế. Bộ bốn field lõi và chính sách link trước là quyết định riêng của Prime OS, không phải quy tắc chung được sao chép từ đối thủ.')
head('Cách dùng tài liệu')
p('Trang 2–4: mô hình và rule. Trang 5–8: flow, UI và case vận hành. Trang 9–10: prototype đã có gì và kỹ thuật cần chốt gì. Trang 11: kế hoạch kiểm chứng. Trang 12: nguồn tham khảo.')

# 2 — Mental model and seller needs.
page('01 Bài toán seller và mô hình sản phẩm')
head('Seller cần hoàn thành công việc gì')
bullet('Đưa các listing đang bán về một nơi để quản lý mà không phải nhập lại dữ liệu có sẵn.')
bullet('Nhận ra listing nào cùng sản phẩm, SKU nào cùng biến thể và dữ liệu nào còn khác nhau giữa các shop.')
bullet('Mở rộng sang shop hoặc sàn mới mà vẫn giữ những nội dung, giá và điều kiện riêng của từng nơi bán.')
bullet('Biết chính xác đang thiếu gì, thao tác nào bị chặn, tác động đến đâu và có thể tiếp tục việc khác hay không.')
head('Bốn lớp dữ liệu cần phân biệt')
table(['Lớp', 'Vai trò trong Prime OS', 'Không nên hiểu là'], [
    ['Product Master', 'Hồ sơ sản phẩm chuẩn nội bộ; chứa dữ liệu dùng chung và một hoặc nhiều SKU.', 'Bản sao nguyên trạng của một listing; hoặc bộ field bắt buộc của tất cả sàn.'],
    ['Master SKU', 'Định danh đơn vị bán hoặc biến thể cụ thể; liên hệ với giá và tồn theo kho.', 'Mã SKU trên các sàn phải giống hệt nhau.'],
    ['Shop listing và Shop SKU', 'Phiên bản bán tại một shop; giữ ID, nội dung, taxonomy và thiết lập riêng.', 'Đã liên kết thì mọi dữ liệu lập tức bị Master ghi đè.'],
    ['Tồn theo kho', 'Số lượng và trạng thái tồn của một SKU tại một nguồn/kho cụ thể.', 'Số lượng hiển thị trên nhiều listing có thể cộng thành tồn vật lý.'],
], [1.28, 3.02, 2.76])
head('Quan hệ cần hiển thị rõ trên UI')
p('Một Master → nhiều Master SKU. Nhiều listing → một Master khi cùng sản phẩm. Mỗi Shop SKU → một Master SKU đã xác nhận trước khi dùng cho đồng bộ giá/tồn. Các shop khác nhau có thể cùng tham chiếu một Master SKU; không vì vậy mà tạo thêm hàng tồn.')
head('Làm giàu dữ liệu không đồng nghĩa ghi đè')
p('Listing là nguồn bằng chứng để đề xuất bổ sung Master. Mỗi thay đổi phải thể hiện nguồn, giá trị cũ và giá trị mới. Thông tin chỉ có ý nghĩa trên một sàn được giữ ở listing hoặc cấu hình theo sàn. Khác category hay SKU code chưa đủ để kết luận khác sản phẩm; ngược lại, tên gần giống hoặc số SKU bằng nhau cũng chưa chứng minh cùng sản phẩm.')
p('Phạm vi nghiên cứu gồm tạo/nhập Master, review listing, field và SKU mapping, readiness, cấu hình listing, giá/tồn, và khả năng sửa sai. Thuế, thanh toán, định tuyến đơn hàng và tích hợp API chi tiết nằm ngoài phạm vi này.')

# 3 — Evidence, not feature parity claims.
page('02 Research và hàm ý thiết kế')
head('Độ đầy đủ phải có ngữ cảnh')
p('Akeneo tính completeness theo family, channel và locale; cùng một sản phẩm có thể đủ dữ liệu ở ngữ cảnh này nhưng thiếu ở ngữ cảnh khác. Hàm ý cho Prime OS: tách readiness lõi của Master khỏi yêu cầu theo listing, không dùng một tiến độ chung để đại diện mọi sàn. [R1]')
head('Yêu cầu của sàn không phải một checklist tĩnh')
p('Amazon Product Type Definitions cung cấp schema theo product type và marketplace, có hỗ trợ ngữ cảnh seller và parentage. Hàm ý: connector cần trả rule có phiên bản và ngữ cảnh; UI không tự suy ra yêu cầu Amazon từ category nội bộ, cũng không áp rule Amazon cho Shopee hoặc Lazada. [R2]')
head('Active và publish là hai khái niệm khác nhau')
p('Shopify mô tả trạng thái Active là sẵn sàng bán và có thể được publish; Active không tự publish lên sales channel. Đây là tiền lệ phù hợp để giải thích trạng thái, không phải khẳng định Prime OS có cùng toàn bộ vòng đời sản phẩm. [R3]')
head('Tồn được quản lý theo nguồn riêng')
p('Shopify theo dõi tồn độc lập tại các location được gán và phân biệt tồn với cấu hình fulfillment. Hàm ý: bảng SKU cần cho seller thấy kho nào có dữ liệu thực tế, không tạo cảm giác SKU đã có hàng ở mọi kho. Thêm kho và bật stock sync phải là hai hành động khác nhau. [R4]')
head('Giảm form nhưng không giấu quyết định quan trọng')
p('NN/g khuyến nghị đưa lựa chọn thường dùng lên trước và chỉ mở chi tiết ít dùng khi cần. Áp dụng: hiển thị ngay dữ liệu lõi, khác biệt, lỗi và cặp SKU; phần giải thích hoặc metadata mở theo nhu cầu. Không giấu mọi comparison field sau một nút điều hướng khác. [R5]')
head('Cho sửa sai nhưng không tạo tác động âm thầm')
p('NN/g nhấn mạnh quay lại, hủy và khôi phục để người dùng giữ quyền kiểm soát. Với Prime OS, unlink là công cụ sửa quan hệ, không phải lời hứa đảo ngược mọi thay đổi đã gửi ra sàn hoặc mọi nghiệp vụ tồn. Phạm vi khôi phục phải được mô tả trước khi xác nhận. [R6]')
head('Thông báo phải nhận biết được và có đường xử lý')
p('W3C yêu cầu lỗi được mô tả bằng chữ và xác định mục bị lỗi; tooltip tùy chỉnh xuất hiện khi hover/focus phải đáp ứng điều kiện tương tác tương ứng. Vì vậy, lý do chặn và hành động sửa phải nằm trực tiếp trên UI; tooltip chỉ bổ sung giải thích. [R7, R8]')
p('Các hàm ý trên là suy luận thiết kế áp dụng vào ngữ cảnh Prime OS. Chưa có dữ liệu để khẳng định giảm bao nhiêu phần trăm thao tác hoặc tăng conversion của seller.')

# 4 — Rules.
page('03 Tách điều kiện liên kết Active và sync')
p('Không dùng một trạng thái “mapping thành công” để bao trùm cả lưu quan hệ, khớp SKU, cập nhật Master và gửi dữ liệu ra sàn. Seller cần nhìn thấy kết quả của từng lớp.')
table(['Tình huống', 'Lưu liên kết', 'Master Active', 'Sync listing'], [
    ['Đủ dữ liệu lõi và cặp SKU đã xác nhận', 'Được sau xác nhận', 'Có thể kích hoạt', 'Chỉ khi đủ rule sàn và có cấu hình sync phù hợp'],
    ['Thiếu SKU, tên, mô tả hoặc giá lõi', 'Có thể lưu nháp / xử lý sau', 'Không kích hoạt mới', 'Không bật sync theo flow này'],
    ['Master đủ lõi nhưng thiếu required của sàn A', 'Được', 'Không bị chặn bởi sàn A', 'Chặn listing A; listing B xét độc lập'],
    ['SKU mapping chưa hoàn tất', 'Được ở trạng thái pending', 'Không tự hạ Active của Master hiện có', 'Chặn giá/tồn của mapping chưa xác nhận'],
    ['Khác brand, mã hoặc taxonomy', 'Cần seller kiểm tra ý nghĩa', 'Không tự kết luận thiếu lõi', 'Xét rule và mapping thực tế'],
    ['Dữ liệu nguồn đổi hoặc listing đã có chủ sở hữu', 'Dừng lưu để tải lại / xử lý xung đột', 'Giữ trạng thái đã lưu', 'Không thực hiện ghi dựa trên dữ liệu cũ'],
], [2.13, 1.4, 1.55, 1.98])
head('Dữ liệu lõi theo quyết định hiện tại')
p('Tên và mô tả phải có nội dung; SKU không rỗng và không trùng trong phạm vi kiểm tra; giá phải hợp lệ. Prototype đang kiểm tra giá lớn hơn 0, và giá/SKU của các biến thể được chọn. Giá bằng 0, quy tắc tiền tệ và phạm vi duy nhất của SKU cần được Product và Engineering chốt riêng trước production.')
head('Đủ dữ liệu chưa có nghĩa đã đồng bộ')
p('UI nên phân biệt: “Đã liên kết”, “Cần hoàn tất SKU”, “Master Active”, “Đủ dữ liệu cho sàn”, “Sync đang tắt” và “Đã sync”. Trạng thái cuối chỉ xuất hiện khi hệ thống nhận kết quả thực tế từ connector; không suy ra từ việc bấm Save.')
head('Ví dụ trình bày với team')
p('Master có đủ bốn nhóm dữ liệu lõi. Listing Shopee thiếu thuộc tính bắt buộc theo category; listing Amazon đã đủ dữ liệu theo schema đang áp dụng. Master vẫn Active. Shopee cần xử lý trước khi sync; Amazon chỉ được sync khi SKU mapping và cấu hình liên quan cũng hợp lệ.')
p('Chặn theo nhóm dữ liệu hay chặn toàn bộ lượt sync của một listing vẫn cần contract kỹ thuật rõ ràng. Tối thiểu, không để SKU chưa xác nhận đi vào cập nhật giá hoặc tồn.')

# 5 — Flow.
page('04 Luồng review từ chọn listing đến kết quả')
head('Bước 1 Chọn đúng ý định')
table(['Ý định của seller', 'Hành vi và ngữ cảnh UI'], [
    ['Review nhiều sản phẩm riêng', '“Review selected listings” đưa vào hàng đợi từng listing với Master riêng; không gộp mặc định.'],
    ['Nhiều listing của cùng sản phẩm', '“Group as one product” → link vào một Master hoặc tạo một Master mới. Luôn hiển thị N listings → 1 Master.'],
    ['Chưa có Master nào', 'Ẩn cột gợi ý Master và CTA chọn Master không khả dụng. Cho tạo riêng từng listing hoặc chủ động nhóm cùng sản phẩm.'],
    ['Đã có Master', 'Giữ cột Product Master và CTA Review suggestion / Choose Master. Gợi ý chưa phải liên kết đã lưu.'],
], [2.16, 4.9])
head('Bước 2 Review quan hệ và dữ liệu cùng một ngữ cảnh')
p('Hiển thị các listing trong hàng đợi và Master đích. Khi nhóm, Master đích áp dụng cho toàn bộ nhóm, không chỉ listing đang xem. Nếu tạo mới, “Nguồn dữ liệu khởi tạo” chỉ quyết định dữ liệu được điền trước; không làm thay đổi số listing sẽ được link.')
p('Seller kiểm tra cùng sản phẩm, nguồn field và các cặp SKU. Các khác biệt được phân loại thành thông tin, cần xác nhận, hoặc thiếu dữ liệu có thể xử lý tại chỗ. Có thể chuyển listing tiếp theo mà không buộc phải đánh dấu listing hiện tại là đã kiểm tra.')
head('Bước 3 Xác nhận theo mức độ sẵn sàng')
bullet('Đã đủ dữ liệu: cho lưu mapping; nếu tạo mới và đủ lõi, CTA nêu rõ “Create, activate & link”. Nếu cập nhật Master cũ, hiển thị chính xác dữ liệu sẽ thay đổi.')
bullet('Chưa muốn hoàn tất: cho “Link now, finish later” hoặc “Save draft & link”. Nêu listing/SKU nào còn pending, phần sync nào chưa được dùng và đường quay lại xử lý.')
bullet('Dữ liệu đã đổi, Master bị archive, thiếu quyền hoặc quan hệ trùng: chặn thao tác ghi không hợp lệ và cung cấp cách tải lại/chọn đích khác. Đây không phải lỗi có thể bỏ qua bằng cảnh báo.')
head('Bước 4 Trả kết quả đủ rõ để seller yên tâm')
p('Màn hình kết quả cần nói rõ số listing đã link, Master đích, trạng thái Active/Draft, số mapping còn pending và “Chưa publish; sync chưa tự bật”. Với xử lý hàng loạt, liệt kê riêng thành công và thất bại; không báo thành công toàn bộ khi chỉ lưu được một phần.')
p('Ví dụ: sổ vẽ và kéo may phải đi qua review riêng. Hai listing cùng mẫu sổ vẽ trên hai shop có thể được nhóm vào một Master, nhưng vẫn phải kiểm tra các biến thể thực sự tương ứng.')

# 6 — Layout and mapping examples.
page('05 Layout review và chỉnh field tại chỗ')
p('Đề xuất một màn hình làm việc tập trung: đọc nhanh trước, chỉnh khi cần. Không bắt seller chuyển sang Product detail chỉ để hoàn thành các field đang được review.')
head('Thứ tự thông tin từ trên xuống dưới')
p('1  Phạm vi thao tác: review riêng hay N listings → 1 Master; tên Master đích và trạng thái chưa lưu.\n2  Listing đang xem và Master đích: tên shop nổi bật, tên sàn ở dòng dưới, logo đúng channel.\n3  Dữ liệu lõi và bảng comparison mở sẵn; khác biệt cần xử lý đặt trước.\n4  SKU mapping và biến thể; hiển thị dữ liệu nguồn đã sync, không bắt nhập lại.\n5  Yêu cầu riêng của từng listing; mô tả, media, shipping và định danh mở chi tiết khi cần.\n6  Footer cố định: tiến độ, tác động của lần lưu và một CTA chính theo trạng thái.')
head('Bảng thể hiện nguồn và kết quả sẽ lưu')
table(['Field Master', 'Nguồn listing', 'Giá trị sẽ lưu', 'Thao tác'], [
    ['Tên sản phẩm *', 'Shop A · Listing title\nSketchbook Pro A5', 'Sketchbook Pro A5', 'Đổi nguồn / sửa'],
    ['Master SKU *', 'Mã nội bộ', 'PRD-0003', 'Sửa mã; kiểm tra trùng'],
    ['Mô tả *', 'Shop A · Description', 'Giấy vẽ 200 gsm…', 'Xem đủ / sửa'],
    ['Giá cơ bản *', 'Shop A · Price\n35 USD', '35 USD', 'Sửa; kiểm tra tiền tệ'],
    ['Brand', 'Shop B · Brand\nPrime Craft', 'Prime Craft', 'Giữ Master / lấy nguồn'],
], [1.18, 2.32, 2.0, 1.56])
p('Ví dụ minh họa. Dấu * là dữ liệu lõi. Với sản phẩm biến thể, mã và giá từng SKU được kiểm tra ở bảng bên dưới; required riêng của sàn được gắn tên sàn, không dùng chung dấu * của Master.')
head('Tương tác đề xuất')
bullet('Click vào giá trị hoặc Edit ngay dòng để sửa; có lưu/hủy thay đổi tại chỗ. Cho đổi field nguồn, chọn listing nguồn khác hoặc nhập thủ công. Bản chụp nguồn không bị sửa theo.')
bullet('Đánh dấu “Giữ Master”, “Từ listing”, “Đã sửa” và giá trị trước/sau. Đổi nguồn khởi tạo không được âm thầm xóa phần seller đã chỉnh; cần preview hoặc xác nhận.')
bullet('Không dùng “Missing data” chung chung: ghi “Listing chưa có GTIN”, “Master thiếu giá” hoặc “Lazada cần Material”. Không biến mọi khác biệt thành lỗi màu vàng.')
head('Sạch nhưng không giấu lỗi')
p('Tooltip dành cho ý nghĩa field, nguồn và cách tính. Lỗi chặn, phạm vi cập nhật và việc sync đang tắt phải hiển thị trực tiếp. Khi lỗi phát sinh, có link đưa đến đúng dòng và focus editor; hỗ trợ bàn phím, không chỉ dùng màu. [R7, R8]')

# 7 — Edge cases.
page('06 Ma trận xử lý SKU và các ngoại lệ')
p('SKU mapping là đối chiếu đơn vị bán thực tế, không chỉ so sánh hai chuỗi mã. Bảng dưới là hành vi mục tiêu để thống nhất triển khai và nghiệm thu; những phần đã có trong prototype được ghi ở trang 9.')
table(['Case', 'Xử lý và thông tin seller cần thấy', 'Khi chưa hoàn tất'], [
    ['1 SKU → 1 SKU', 'Gợi ý cặp tương ứng; seller kiểm tra sản phẩm và pack. Khác mã không tự chặn.', 'Cho lưu pending nếu chưa xác nhận.'],
    ['3 SKU → Master 1 SKU', 'CTA Set up SKUs; đề xuất biến thể từ nguồn có thật, hiển thị việc chuyển cấu trúc Master.', 'Link trước; chặn giá/tồn của mapping chưa xác nhận.'],
    ['3 SKU → Master 3 SKU', 'Đối chiếu từng cặp theo option, mã và đơn vị bán. Cùng số lượng chưa đủ.', 'Nêu rõ cặp còn thiếu; không ép chọn theo thứ tự.'],
    ['3 SKU → Master 5 SKU', 'Map ba SKU đang có; không yêu cầu mọi listing phải phủ hết SKU của Master.', 'Hai SKU Master còn lại được giữ nguyên.'],
    ['Nhiều shop có cùng biến thể', 'Các Shop SKU tương ứng có thể map cùng Master SKU giữa các listing.', 'Không nhân bản tồn vì có thêm shop.'],
    ['Hai SKU trong một listing cùng trỏ một SKU', 'Không tự gộp các biến thể khác nhau. Rule ngoại lệ cho alias/bundle phải được chốt riêng.', 'Giữ pending; yêu cầu sửa cặp hoặc xác minh.'],
    ['Biết có 3 SKU nhưng thiếu bản ghi con', 'Nêu “Đã tải 1/3 SKU”; cho tải lại hoặc sync lại nguồn. Không tự bịa mã con.', 'Cho lưu quan hệ nếu hợp lệ; chưa xác nhận mapping.'],
    ['SKU trùng hoặc chưa biết cấu trúc', 'Chỉ đúng dòng trùng hoặc yêu cầu xác minh Single/Variants. Có CTA xử lý tại chỗ.', 'Không dùng dữ liệu mơ hồ để cập nhật giá/tồn.'],
], [1.58, 3.47, 2.01])
head('Case trong ảnh cuối của luồng review')
p('Nhóm đang chứa sổ vẽ, bộ cọ và tranh nhưng Master đích là bộ cọ. Có hai vấn đề độc lập: ý định nhóm có thể sai và listing sổ vẽ có ba SKU trong khi Master là single. UI phải nêu cả hai, cho quay lại review riêng/đổi nhóm, và không chỉ vô hiệu hóa nút Continue.')
head('Ranh giới an toàn')
p('Có thể cho lưu quan hệ pending khi seller chủ động lựa chọn; không vì vậy mà tự xác nhận cùng sản phẩm. Thay Master đích hoặc sửa cặp SKU phải làm mất hiệu lực xác nhận liên quan. Tạo biến thể không tự chia tồn cũ, không tự đổi pack và không tự bật sync.')

# 8 — Other product functionality.
page('07 Giá tồn kho và vận hành sau mapping')
head('Seller không cần chọn một con số tồn để copy từ listing')
p('Số tồn hiển thị trên sàn là dữ liệu báo về của listing, chưa đủ để kết luận đó là tồn vật lý độc lập. Hai sàn có thể cùng dùng một kho. Cộng hai số này vào Master sẽ có nguy cơ đếm trùng. Mô hình tham chiếu theo location giúp phân biệt nguồn tồn và cách fulfillment. [R4]')
bullet('Chỉ hiển thị kho có quan hệ hoặc dữ liệu liên quan đến SKU; các kho chưa gán nằm trong “Add stock location”. Không trình bày năm kho “Not recorded” như năm nơi đã có hàng.')
bullet('Tách “Recorded stock” khỏi “Available to sell”. Phân biệt 0 đã xác nhận, chưa ghi nhận và chưa kiểm chứng; không dùng số 0 để thay cho thiếu dữ liệu.')
bullet('Kho seller quản lý có quyền ghi nhận/điều chỉnh; tồn do channel hoặc đơn vị fulfillment quản lý được tách và ghi Read only theo khả năng connector.')
bullet('SKU mới cần được lưu trước khi thiết lập tồn. Gán kho không đồng nghĩa phân bổ tồn cho listing và không tự bật stock sync.')
head('Giá và dữ liệu riêng theo shop')
p('Giá Master là dữ liệu chuẩn nội bộ; giá bán theo shop có thể chịu currency, chính sách giá và override. Không copy số tiền khác currency mà bỏ đơn vị; không gọi đó là quy đổi nếu chưa có tỷ giá và quy tắc làm tròn. Việc lấy dữ liệu vào Master và việc áp Master ra listing phải là hai quyết định riêng.')
head('Tạo listing mới và link listing có sẵn')
p('Giữ hai CTA khác nhau. “Create channel listings” chọn shop để tạo bản ghi mới; “Link existing listings” chọn listing đã sync nhưng chưa liên kết. Shop đã liên kết cần hiển thị trạng thái hoặc bị loại khỏi luồng tạo theo policy hiện tại; bỏ tick ở bước chọn shop không được hiểu là unlink listing đã tồn tại.')
head('Thông tin trên Product detail')
p('Overview ưu tiên trạng thái, số listing và việc cần xử lý. Product data chứa dữ liệu chuẩn và media gọn. Pricing & Inventory thể hiện giá, SKU và nguồn tồn. Channel listings thể hiện điều kiện theo shop. Readiness dùng cùng vị trí/cột và spacing giữa các tab; không lặp banner thiếu dữ liệu cùng một mục ở nhiều nơi.')
head('Unlink và khôi phục')
p('Confirmation phải nêu listing nào bị gỡ quan hệ, sync/config nào bị ngắt và dữ liệu Master nào được giữ. Unlink không xóa listing trên sàn, không tự hoàn tác thay đổi đã publish và không đảo ngược đơn hàng/tồn. Audit log hỗ trợ truy vết; khả năng rollback từng nghiệp vụ vẫn cần contract riêng. [R6]')

# 9 — Implementation audit.
page('08 Prototype hiện tại và giới hạn triển khai')
p('Đối chiếu mã nguồn cục bộ ngày 09/10/2026. Các mục dưới xác nhận logic thể hiện trong prototype, không xác nhận tính năng đã chạy với dữ liệu thật trên từng marketplace.')
table(['Năng lực', 'Bằng chứng trong prototype', 'Giới hạn cần biết'], [
    ['Review riêng hoặc nhóm', 'ProductListingIntake có queue riêng và Group as one product.', 'Cần đo seller có hiểu khác biệt giữa hai ý định.'],
    ['Readiness lõi', 'product-master-readiness kiểm tra tên, mô tả, SKU và giá.', 'Chính sách giá 0, SKU uniqueness và validation phía server chưa được chốt ở tài liệu này.'],
    ['Rule theo listing', 'listing-requirements có channel, category, revision và origin; listing-sync-readiness đánh giá từng listing.', 'Có schema prototype; chưa đủ bằng chứng connector cập nhật rule thật, rule điều kiện và thời hạn hiệu lực.'],
    ['Field mapping có nguồn', 'listing-field-mapping lưu source, field, giá trị và transform; hỗ trợ quyết định thủ công.', 'Gợi ý đang dựa trên quy tắc; không coi là bằng chứng đã có AI production.'],
    ['SKU mapping và recovery', 'Gợi ý mã/biến thể không mơ hồ; phân biệt thiếu nguồn, thiếu SKU Master và trùng cặp.', 'Không tự gộp nhiều source SKU vào một SKU; ngoại lệ cần thiết kế riêng.'],
    ['Link xử lý sau', 'confirmListingIntake có defer và review_pending; tách draft khỏi dữ liệu đã xác nhận.', 'Cần API bảo đảm transaction, retry và xử lý lỗi nhiều listing ở backend.'],
    ['Kho và unlink', 'VariantInventoryDetails phân nhóm kho; unlinkListingFromMaster kiểm tra snapshot và gỡ quan hệ cục bộ.', 'Không chứng minh rollback đầy đủ trên sàn hoặc tích hợp WMS.'],
], [1.43, 3.12, 2.51])
head('Không đưa thành cam kết khi trình bày')
p('Chưa có bằng chứng trong phần mã được rà soát rằng MappingEngine/AI tự giải quyết mọi field/value, tự sinh SKU đúng mọi trường hợp, chống race-condition xuyên hệ thống hoặc khôi phục mọi tác động đã sync. Đây là yêu cầu kỹ thuật cần xác nhận, không phải tính năng đã hoàn thiện.')
p('Product store hiện có lưu trữ localStorage. Việc kiểm tra snapshot/version trong prototype là cơ sở mô phỏng hành vi, không thay thế khóa phiên bản, phân quyền và audit bền vững ở server. Mã trạng thái nội bộ “published” đang được dùng cho Master Active; không được hiển thị như đã đăng bán lên sàn.')
p('Đường dẫn bằng chứng: apps/web/src/lib/ và apps/web/src/components/products/ với tên file trong bảng. Cần đối chiếu lại theo phiên bản release trước khi nghiệm thu production.')

# 10 — Decisions.
page('09 Những điểm cần chốt với kỹ thuật')
p('Ưu tiên thống nhất contract đầu vào, đầu ra và phạm vi ảnh hưởng trước khi tiếp tục tinh chỉnh copy trên UI. Những câu hỏi sau quyết định trực tiếp seller có bị chặn đúng chỗ hay không.')
table(['Cần chốt', 'Đề xuất làm cơ sở thảo luận', 'Nhóm phụ trách'], [
    ['Required theo sàn lấy từ đâu', 'Schema theo shop/marketplace/category/product type, locale khi liên quan; có revision, nguồn và cách phát hiện hết hiệu lực.', 'Connector + Product'],
    ['Điều kiện Active', 'Chốt nghĩa hợp lệ của SKU/title/description/price; có cho giá 0 không; xử lý Master Active bị sửa thiếu dữ liệu ra sao.', 'Product + Catalog'],
    ['Pending chặn sync đến mức nào', 'Giá/tồn luôn cần SKU đã xác nhận. Quyết định content/media có được sync độc lập hay toàn listing phải chờ.', 'Sync + Product'],
    ['AI được quyền làm gì', 'Đề xuất, giải thích nguồn, độ chắc chắn; chỉ tự áp theo chính sách được duyệt. Không dùng confidence làm bằng chứng duy nhất.', 'AI/Mapping + Design'],
    ['3 SKU vào Master single', 'Có proposal thêm/chuyển biến thể, preview tác động lên listing đã link; không sửa tồn hoặc mất SKU cũ.', 'Catalog + Inventory'],
    ['Ghi đè và xung đột', 'Preview diff, field ownership, phiên bản nguồn/Master; kiểm tra lại lúc commit; bảo toàn quyết định của seller.', 'Backend + Design'],
    ['Lưu hàng loạt và khôi phục', 'Chọn atomic theo nhóm hoặc partial success có receipt từng listing; retry idempotent; xác định phạm vi undo/unlink.', 'Backend + QA'],
], [1.6, 4.07, 1.39])
head('Contract readiness tối thiểu đề xuất')
p('Mỗi issue nên trả về đối tượng bị ảnh hưởng, field hoặc SKU, loại issue, hành động bị chặn, lý do và hành động khắc phục; kèm rule revision và thời điểm kiểm tra. UI dùng chung dữ liệu này ở review, readiness và cấu hình sync để không hiển thị ba kết luận khác nhau.')
head('Điểm phản biện cần giữ')
p('“Có thể unlink” không đủ để bỏ kiểm soát trước sync, vì tác động ngoài hệ thống có thể không đảo ngược được. Ngược lại, “Master là hồ sơ gốc duy nhất” không có nghĩa bắt seller hoàn thiện mọi dữ liệu của mọi sàn trước khi link. Giải pháp cân bằng là cho lưu quan hệ có trạng thái, nhưng kiểm soát chặt việc dùng dữ liệu và gửi thay đổi ra ngoài.')

# 11 — Validation.
page('10 Ưu tiên triển khai và kiểm chứng với seller')
head('Thứ tự ưu tiên')
bullet('P0 — Đúng logic và an toàn: tách link/Active/sync; tách review riêng/nhóm; SKU recovery; lưu pending; kiểm tra dữ liệu đổi; không tự publish hoặc sửa tồn.')
bullet('P1 — Giảm công sức: prefill từ nguồn, bảng mapping sửa tại chỗ, giữ state khi quay lại, receipt từng listing, readiness theo sàn và quản lý pending có lối quay lại.')
bullet('P2 — Tăng hiệu suất sau khi đo: quick review có điều kiện, bulk edit có phạm vi rõ, gợi ý thông minh và tái sử dụng mapping đã được phê duyệt.')
head('Kiểm thử định tính đề xuất')
p('Vòng đầu với 5–8 seller hoặc nhân sự vận hành đại diện cho single SKU, nhiều biến thể và nhiều shop; đây là mẫu thăm dò, không dùng để suy rộng tỷ lệ thống kê. Quan sát thao tác và yêu cầu người tham gia diễn giải kết quả, không hướng dẫn trước ý nghĩa của các nút.')
table(['Nhiệm vụ', 'Bằng chứng đạt kỳ vọng'], [
    ['Chọn hai sản phẩm khác nhau để review', 'Không vô tình tạo một Master chung; hiểu Master đích của từng listing.'],
    ['Hai shop cùng sản phẩm nhưng khác mã SKU', 'Biết đổi nguồn field và xác nhận từng biến thể, không sửa dữ liệu nguồn để ép khớp.'],
    ['3 SKU vào Master single', 'Tìm được Set up SKUs hoặc Link later; hiểu phần sync nào chưa được phép.'],
    ['Thiếu required riêng của một sàn', 'Không nghĩ Master hoặc mọi sàn đều bị lỗi; tìm đúng field và đúng listing để sửa.'],
    ['SKU mới chưa có tồn', 'Phân biệt chưa ghi nhận với 0; gán đúng kho mà không copy tổng tồn từ các shop.'],
    ['Rời màn hình rồi quay lại / unlink', 'Biết phần nào đã lưu; giữ được dữ liệu hợp lệ và hiểu phạm vi gỡ liên kết.'],
], [2.47, 4.59])
head('Đo gì sau khi triển khai')
p('Đo thời gian đến kết quả lưu, tỷ lệ hoàn thành không cần trợ giúp, số lần quay lại sai bước, tỷ lệ sửa mapping đã xác nhận và tuổi của pending. Dùng sync sai đích, ghi đè ngoài ý muốn và tồn bị đếm trùng làm chỉ số bảo vệ.')
p('Ghi sự kiện chọn mode, sửa field/SKU, defer, confirm và kết quả. Chốt baseline qua test/pilot trước khi đặt mục tiêu; không dùng số ước đoán như kết quả nghiên cứu.')
head('Tiêu chí nghiệm thu xuyên suốt')
p('Lỗi có cách xử lý; chuyển listing không tự xác nhận; checked chưa phải saved; thiếu required sàn A không chặn B; retry không tạo link trùng. Tạo/link/Active không tự publish hoặc bật sync; cấu hình đã có được giữ theo phạm vi đã xác nhận.')

# 12 — Reference trail.
page('11 Nguồn tham khảo và cách diễn giải')
p('Nguồn được truy cập ngày 09/10/2026. Các liên kết dưới có thể mở trực tiếp trong Word. Tài liệu benchmark dùng để chứng minh nguyên tắc và khả năng có thật ở hệ thống tham chiếu, không phải cam kết tương đương tính năng hoặc hiệu năng của Prime OS.')

refs = [
    ('R1', 'Akeneo — Follow your product’s completeness', 'https://help.akeneo.com/187-serenity-follow-your-products-completeness?kb_language=en_US', 'Cơ sở cho completeness theo family, channel, locale và chỉ ra thuộc tính còn thiếu.'),
    ('R2', 'Amazon — Retrieve a Product Type Definition', 'https://developer-docs.amazon.com/sp-api/lang-en_EN/docs/retrieve-a-product-type-definition', 'Cơ sở cho requirements có ngữ cảnh product type, marketplace, seller và parentage.'),
    ('R3', 'Shopify — ProductStatus', 'https://shopify.dev/docs/api/admin-graphql/latest/enums/ProductStatus', 'Cơ sở phân biệt Active với việc publish lên sales channel.'),
    ('R4', 'Shopify — Inventory management for multiple locations and apps', 'https://help.shopify.com/en/manual/inventory-and-locations/setup/multi-managed-inventory', 'Cơ sở phân biệt tồn tại location và việc dùng location để fulfillment.'),
    ('R5', 'Nielsen Norman Group — Progressive Disclosure', 'https://www.nngroup.com/articles/progressive-disclosure/', 'Cơ sở ưu tiên thông tin thường dùng và mở chi tiết theo nhu cầu; không giấu thông tin thiết yếu.'),
    ('R6', 'Nielsen Norman Group — User Control and Freedom', 'https://www.nngroup.com/articles/user-control-and-freedom/', 'Cơ sở cho quay lại, hủy, sửa sai và bảo toàn quyền kiểm soát của người dùng.'),
    ('R7', 'W3C — Understanding SC 3.3.1 Error Identification', 'https://www.w3.org/WAI/WCAG22/Understanding/error-identification.html', 'Cơ sở mô tả lỗi bằng chữ và xác định đối tượng bị lỗi; không chỉ dùng màu hoặc nút disabled.'),
    ('R8', 'W3C — Understanding SC 1.4.13 Content on Hover or Focus', 'https://www.w3.org/WAI/WCAG22/Understanding/content-on-hover-or-focus.html', 'Cơ sở tương tác tooltip: có thể đóng, di chuyển con trỏ vào và giữ hiển thị theo điều kiện phù hợp.'),
]
for rid, label, url, use in refs:
    para = p('[' + rid + '] ')
    para.paragraph_format.space_after = Pt(2)
    para.paragraph_format.keep_with_next = True
    link(para, label, url)
    para = p(use)
    para.paragraph_format.space_after = Pt(9)
    for r in para.runs:
        r.font.size = Pt(10.5)

head('Độ tin cậy của các loại kết luận')
p('Quyết định sản phẩm: yêu cầu đã thống nhất trong hội thoại, cần được team phê duyệt thành spec. Quan sát prototype: hành vi có bằng chứng mã nguồn tại thời điểm rà soát. Research tham chiếu: thông tin trong nguồn chính thức. Đề xuất: thiết kế hoặc contract cần kiểm chứng thêm bằng usability test và triển khai backend.')
p('Thông điệp chốt: seller có thể bắt đầu quản lý trước khi dữ liệu hoàn hảo, nhưng hệ thống phải nói rõ dữ liệu nào đã được xác nhận, đang được dùng ở đâu và điều gì chưa được phép đồng bộ.')

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
print(f'Paragraphs: {len(doc.paragraphs)}; tables: {len(doc.tables)}; references: {len(refs)}')
